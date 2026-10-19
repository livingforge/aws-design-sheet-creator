"""Lossless text-first XLSX renderer for normalized design documents."""
import json
import math

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .document_layout import LAYOUT, LAYOUT_VERSION
from .excel_report import HEADER_FILL, HEADER_FONT


EDIT_COLUMNS=('修正案（値）','修正理由','適用先（共通/環境）')
IMPORT_KEY='@import-key'


def text(value):
    return value if isinstance(value,str) else json.dumps(value,ensure_ascii=False,sort_keys=True)


def chunks(value,limit=120):
    """Split visible text; no Excel 32767-character truncation or hidden tall rows."""
    raw=ILLEGAL_CHARACTERS_RE.sub(lambda m:'\\u%04x'%ord(m[0]),text(value))
    result=[];part='';lines=0
    for char in raw:
        part+=char;lines+=char=='\n'
        if len(part)>=limit or lines>=6:
            result.append(part);part='';lines=0
    if part or not result:result.append(part)
    return result


def setting(item):
    states={'DECIDED':'決定済み','UNDECIDED':'未決定','NOT_APPLICABLE':'適用外'}
    modes={'LITERAL':'明示値','DEFAULT':'既定動作（値未解決）','OMITTED':'明示省略','RESOURCE_REF':'内部参照',
           'EXTERNAL_REF':'外部参照','EXPRESSION':'未解決式','SECRET_REF':'秘密参照','COMPOSITE':'複合設定（内部の決定状態を保持）'}
    header=states[item.status]+' / '+item.classification
    if item.mode:header+=' / '+modes[item.mode]
    payload=item.value if item.mode in ('LITERAL','COMPOSITE') else item.expression if item.mode=='EXPRESSION' else item.ref
    return header+('\n'+text(payload) if payload is not None or item.mode=='LITERAL' else '')


def observation(item):
    states={'NO_CFN':'CFNなし','NO_BINDING':'対応付けなし','RESOURCE_MISSING':'リソースなし',
            'TYPE_MISMATCH':'型の不一致','ABSENT':'定義なし','LITERAL':'定義値',
            'EXPRESSION':'未解決式','REDACTED':'秘密の実値は非表示','CONDITIONAL':'条件・マクロ未解決'}
    return states[item.state]+('\n'+text(item.value) if item.value is not None else '')


def workbook(document, *, services=None, service_only=False):
    book=Workbook();book.remove(book.active)
    layout=dict(LAYOUT)
    if document.model_version=='1.1' and not service_only:
        layout['地域別資源']=(['資源ID','リージョン','公式スキーマSHA-256','公式取得元'],[36,24,70,100])
        if document.review_report:
            layout['原資料状態']=(['原資料設定ID','区分','元の状態','元の検証状態','生成項目ID','原資料位置'],[24,15,36,30,60,55])
            layout['実行サマリ根拠']=(['記録ID','区分','JSON'],[32,30,120])
        if getattr(document,'decision_trace',{}).get('status','NOT_CHECKED')!='NOT_CHECKED':
            layout['設計方針']=(['方針ID','章','題名・状態','方針','理由・代替案','要件・検査','適用範囲・具体化','出典'],[30,26,34,65,55,40,48,26])
            layout['方針トレース']=(['方針ID','設定ID・パス','設定の状態','値の解決','入力追跡'],[30,60,30,22,70])
            layout['方針指摘']=(['指摘ID','区分・コード','方針・章・要件','設定・資源','内容'],[60,40,34,50,70])
            if document.decision_trace.get('expansions'):
                layout['方針展開']=(['展開ID','方針ID','設定ID','結果','展開値','入力追跡'],[30,30,45,24,60,40])
    if not service_only:
        layout['必要入力']=(['検査ID','項目ID・パス','必須種別・必要性','条件・根拠','入力状態・解決','結果','対応方法'],[36,42,28,70,26,25,58])
        layout['必要入力根拠']=(['根拠ID','区分','内容'],[36,26,120])
        if document.evaluation.status!='NOT_EVALUATED':
            layout['AWS評価']=(['検査ID','規則・資源・パス','判定','理由','入力追跡'],[36,48,24,85,65])
            layout['AWS未検査']=(['検査ID','種別・資源・パス','理由','入力追跡'],[36,52,90,70])
            layout['AWS評価根拠']=(['根拠ID','区分','内容'],[36,26,120])
        if document.design.project_policies:
            layout['案件方針']=(['検査ID','要件・宣言ID','資源・設定','条件・期待値','結果','理由・例外','入力追跡'],[36,35,45,55,22,55,58])
            layout['案件方針根拠']=(['根拠ID','区分','内容'],[36,26,120])
            layout['案件要件範囲']=(['要件ID','関連する宣言ID','登録状態','評価範囲'],[32,70,30,90])
    if document.model_version=='1.1':
        selected=sorted(set(document.resource_services.values()) if services is None else services)
        layout={k:v for k,v in layout.items() if not k.endswith('設定') and (not service_only or k=='概要')}
        columns,widths=LAYOUT['EC2設定']
        # Reviewers type proposals only into these columns; aws-design-changes reads nothing else.
        for service in selected:layout[service+'設定']=([*columns,*EDIT_COLUMNS],[*widths,30,40,18])
    sheets={}
    for name,(columns,widths) in layout.items():
        ws=book.create_sheet(name);sheets[name]=ws
        ws.append([*columns,'分割行'])
        for cell in ws[1]:cell.fill=HEADER_FILL;cell.font=HEADER_FONT
        for i,width in enumerate([*widths,10],1):ws.column_dimensions[get_column_letter(i)].width=width
        ws.freeze_panes='A2';ws.print_title_rows='1:1';ws.sheet_view.zoomScale=85
        if document.model_version=='1.1':ws.row_dimensions[1].height=24
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.page_margins.left=0.25;ws.page_margins.right=0.25
        ws.oddFooter.center.text='&A — &P / &N'
    stripes={name:0 for name in layout}
    def add(name,values):
        if name not in sheets:return
        ws=sheets[name];parts=[chunks(v,240 if name in ('必要入力根拠','AWS評価根拠','案件方針根拠') else 120) for v in values];count=max(map(len,parts))
        stripes[name]+=1
        for index in range(count):
            if ws.max_row>=1_048_575:raise ValueError('workbook row limit exceeded')
            # Keep identity on every continuation row; other columns concatenate losslessly.
            row=[values[0],*[p[index] if index<len(p) else '' for p in parts[1:]],f'{index+1}/{count}']
            ws.append(row)
            height=1
            for c in ws[ws.max_row]:
                c.data_type='s';c.font=Font(name='Yu Gothic',size=10)
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if stripes[name]%2:c.fill=PatternFill('solid',fgColor='F0F5FA')
                if index==count-1:c.border=Border(bottom=Side(style='hair',color='CAD5E0'))
                if c.value in ('DIFFERENT','UNRESOLVED','MATCH','EXPRESSION_MATCH','REFERENCE_MATCH'):
                    c.fill=PatternFill('solid',fgColor={'DIFFERENT':'FCE4D6','UNRESOLVED':'FFF2CC','MATCH':'E2F0D9','EXPRESSION_MATCH':'FFF2CC','REFERENCE_MATCH':'FFF2CC'}[c.value])
                    c.font=Font(name='Yu Gothic',size=10,bold=True)
                height=max(height,sum(max(1,math.ceil(len(line)/(ws.column_dimensions[c.column_letter].width/2))) for line in str(c.value).split('\n')))
            # Yu Gothic wrapping and Excel's print scaling need vertical leading.
            # Keep the released v1 layout unchanged; pad schema-document rows.
            points=17*(height+1) if document.model_version=='1.1' else 15*height
            ws.row_dimensions[ws.max_row].height=min(390,max(30,points))
    d=document.design;env=next(e for e in d.environments if e.id==document.environment_id)
    if document.model_version=='1.1' and document.review_report and not service_only:
        add('実行サマリ根拠',['@review-report','JSON',document.review_report])
        for row in document.review_report['source_records']:
            add('原資料状態',[row['setting_id'],row['kind'],row['status'],row['verification_status'],row['item_ids'],row['source_pointer']])
    for key,label,value in [('system','システム',d.system.name),('description','説明',d.system.description),
        ('environment','環境',env.name+' / '+env.id),('region','リージョン',env.region),('environment_note','環境説明',env.description),
        ('version','入力 / 共通モデル / レイアウト版',d.schema_version+' / '+document.model_version+' / '+LAYOUT_VERSION),
        ('status','生成 / 設計評価',document.generation_status+' / '+document.evaluation_status),
        ('note','注意','生成成功は構造検証と文書生成の完了を示します。設計評価の合格ではありません。')]:add('概要',[d.system.id+'.'+key,label,value])
    if document.model_version=='1.1':
        add('概要',[IMPORT_KEY,'修正案の取込キー',{'format_version':'1.0','input_sha256':document.input_sha256,
            'environment_id':document.environment_id,'edit_columns':list(EDIT_COLUMNS)}])
    for r in d.resources:
        if document.model_version=='1.1' and not service_only:
            region=document.resource_regions.get(r.id,env.region)
            origin=document.regional_catalog_provenance.get(region,document.catalog_provenance)
            add('地域別資源',[r.id,region,origin['zip_sha256'],origin['source']])
        state=document.resource_comparisons[r.id]
        add('構成一覧',[r.id,r.name,r.type,r.cfn_logical_id or '未対応付け',state,r.rationale,', '.join(r.sources)])
        if state!='PRESENT':add('差異',[r.id,state,r.type,r.cfn_logical_id or '',r.rationale])
    for note in d.requirements:add('要件',[note.id,note.title,note.status,note.text,', '.join(note.resources),', '.join(note.sources)])
    for note in d.operations:add('運用復旧',[note.id,note.title,note.status,note.text,', '.join(note.resources),', '.join(note.sources)])
    for note in d.issues:add('未決事項',[note.id,note.title,note.status,note.text,', '.join(note.resources),', '.join(note.sources)])
    for note in [*d.requirements,*d.operations]:
        if note.status=='UNDECIDED':add('未決事項',[note.id,note.title,note.status,note.text,', '.join(note.resources),', '.join(note.sources)])
    for item in document.items:
        defined=observation(item.defined)
        expected=setting(item.expected)
        if document.model_version=='1.1':
            expected='値の解決: '+item.expected_resolution+'\n'+expected
            defined='値の解決: '+item.defined_resolution+'\n'+defined
        provenance=item.expected.rationale+'\n出典: '+', '.join(item.expected.sources)+'\n入力: '+item.input_pointer
        if item.defined.pointer:provenance+='\nCFN: '+item.defined.pointer
        if getattr(item,'comparison_details',None):
            provenance+='\n比較箇所: '+text([d for d in item.comparison_details if d['result']!='MATCH'])
        if document.model_version=='1.1':provenance+='\n比較根拠: '+item.comparison_basis
        if getattr(item.expected,'decision_ids',None):provenance+='\n方針: '+', '.join(item.expected.decision_ids)
        if item.origin=='decision':provenance+='\n出所: 方針から展開（入力に個別の設定なし）'
        add(item.service+'設定',[item.item_id,item.resource_id,item.name,expected,defined,item.comparison,provenance,
                                 *(['']*len(EDIT_COLUMNS) if document.model_version=='1.1' else [])])
        if item.expected.status=='UNDECIDED':add('未決事項',[item.item_id,item.name,item.expected.status,item.expected.rationale,item.resource_id,', '.join(item.expected.sources)])
        if item.comparison!='MATCH':add('差異',[item.item_id,item.comparison,setting(item.expected),defined,'未比較・未解決は差異なしを意味しません。'+provenance])
        if item.origin=='environment':add('環境差分',[item.item_id,item.resource_id+' / '+item.name,setting(item.common) if item.common else '共通設定なし',setting(item.expected),provenance])
    for conn in d.connections:
        add('接続',[conn.id,conn.source+' → '+conn.target,conn.kind+' / '+conn.status,conn.protocol+' / '+conn.ports,conn.description+'\n'+conn.rationale,document.connection_comparisons[conn.id],', '.join(conn.sources)])
        if document.connection_comparisons[conn.id]!='MATCH':add('差異',[conn.id,document.connection_comparisons[conn.id],conn.source+' → '+conn.target,conn.property or '疎通・依存関係の評価は対象外','接続意図の確認。自動的な接続可否判定は行いません。'])
        if conn.status=='UNDECIDED':add('未決事項',[conn.id,'接続の確認',conn.status,conn.description+'\n'+conn.rationale,conn.source+' → '+conn.target,', '.join(conn.sources)])
    for i,notice in enumerate(document.notices):add('差異',[f'@notice-{i+1}','確認範囲','','',notice])
    for s in d.sources:add('出典',[s.id,s.title,s.locator])
    add('出典',['@input-file','入力ファイル SHA-256',document.input_sha256])
    if document.cfn_sha256:add('出典',['@cfn-file','CFNファイル SHA-256',document.cfn_sha256])
    add('概要',['@input-completeness','必要入力の充足',document.checklist.status])
    add('概要',['@input-denominator','必要入力の集計',document.checklist.summary])
    add('概要',['@input-scope','必要入力の検査範囲',document.checklist.limitations])
    policy=document.project_policy
    add('概要',['@project-policy','案件方針の実行 / 判定',policy.status+' / '+policy.assessment])
    if document.design.project_policies:
        add('概要',['@project-policy-counts','案件方針の集計',policy.summary])
        add('概要',['@project-policy-limits','案件方針の限界',policy.limitations])
        if not service_only:
            add('案件方針根拠',['@project-policy','評価メタデータ',text(policy.model_dump(mode='json',exclude={'rows'}))])
            for row in policy.rows:
                detail_row=sheets['案件方針根拠'].max_row+1
                add('案件方針根拠',[row['policy_check_id'],'案件条件・入力・出典',text(row)])
                start=sheets['案件方針'].max_row+1
                exception=row['exception']
                add('案件方針',[row['policy_check_id'],row['requirement_id']+'\n'+row['policy_id'],
                    row['resource_id']+'\n'+row['path'],row['condition_result']+' / '+row['operator']+' '+text(row['expected']),
                    row['result'],row['reason']+('\n例外: '+exception['id']+' / '+exception['reason'] if exception else ''),
                    row['trace']['row_id']+'\n'+row['trace']['input_pointer']+'\n'+row['trace']['input_state']])
                for index in range(start,sheets['案件方針'].max_row+1):
                    sheets['案件方針'].cell(index,1).hyperlink=f"#'案件方針根拠'!A{detail_row}"
            for row in policy.requirement_coverage:
                add('案件要件範囲',[row['requirement_id'],', '.join(row['policy_ids']),row['state'],
                    '関連付いた宣言だけを検査。要件全体の意味的網羅性は未評価。' if row['policy_ids'] else '案件ルール未登録。この要件は未評価。'])
    if document.evaluation.status!='NOT_EVALUATED':
        evaluation=document.evaluation
        add('概要',['@aws-status','AWS静的評価の実行 / 判定',evaluation.status+' / '+evaluation.assessment])
        add('概要',['@aws-summary','AWS静的評価の集計',evaluation.summary])
        add('概要',['@aws-limits','AWS静的評価の限界',evaluation.limitations])
        if not service_only:
            metadata=evaluation.model_dump(mode='json',exclude={'results','coverage'})
            add('AWS評価根拠',['@evaluation','評価メタデータ',text(metadata)])
            for category,records in [('AWS評価',evaluation.results),('AWS未検査',evaluation.coverage)]:
                for row in records:
                    ident=row.get('evaluation_id','@unmapped')
                    detail_row=sheets['AWS評価根拠'].max_row+1
                    add('AWS評価根拠',[ident,category,text(row)])
                    location=(row.get('rule_id') or row.get('kind') or '')+'\n'+(row.get('resource_id') or '')+'\n'+(row.get('path') or '')
                    traces='\n'.join(dict.fromkeys(m['item_id']+' / '+m['input_pointer'] for m in row.get('trace',[])))
                    reason=row.get('reason','')+'\n'+row.get('bridge_reason','')
                    values=[ident,location,row['verdict'],reason,traces] if category=='AWS評価' else [ident,location,reason,traces]
                    first=sheets[category].max_row+1;add(category,values)
                    for index in range(first,sheets[category].max_row+1):
                        sheets[category].cell(index,1).hyperlink=f"#'AWS評価根拠'!A{detail_row}"
    checklist_evidence={}
    for row in document.checklist.rows if not service_only else []:
        evidence_refs={}
        for key in ('evidence','input_sources'):
            refs=[]
            for entry in row[key]:
                encoded=text(entry)
                if encoded not in checklist_evidence:
                    ident='evidence-'+str(len(checklist_evidence)+1)
                    checklist_evidence[encoded]=ident
                    add('必要入力根拠',[ident,'必要入力の根拠・出典',encoded])
                refs.append(checklist_evidence[encoded])
            evidence_refs[key]=refs
        detail={**{k:row[k] for k in ('condition_result','condition','schema_pointer','requirement_id','requirement_confirmation','input_pointer','input_pointer_exists','default_evidence','detail')},**evidence_refs}
        detail_row=sheets['必要入力根拠'].max_row+1
        add('必要入力根拠',[row['checklist_id'],'条件・出典・入力追跡',text(detail)])
        start=sheets['必要入力'].max_row+1
        add('必要入力',[row['checklist_id'],row['item_id']+'\n'+row['path'],row['requirement_kind']+' / '+row['necessity'],
            row['condition_result']+' / '+row['condition']+'\n根拠: '+', '.join(evidence_refs['evidence'])+'\n詳細は検査IDのリンク先',
            row['input_state']+' / '+row['value_resolution'],row['result'],row['action']])
        for index in range(start,sheets['必要入力'].max_row+1):
            sheets['必要入力'].cell(index,1).hyperlink=f"#'必要入力根拠'!A{detail_row}"
    trace=getattr(document,'decision_trace',{})
    if document.model_version=='1.1' and trace:
        add('概要',['@design-decisions','方針トレースの実行 / 判定',trace['status']+' / '+trace['assessment']])
        if trace['status']!='NOT_CHECKED':add('概要',['@design-decision-counts','方針トレースの集計',trace['summary']])
        add('概要',['@design-decision-limits','方針トレースの限界',trace['limitations']])
    if document.model_version=='1.1' and not service_only and trace.get('status','NOT_CHECKED')!='NOT_CHECKED':
        titles={s['section_id']:s['title'] for s in trace['sections']};rows={r['decision_id']:r for r in trace['decisions']}
        for decision in d.decisions:
            row=rows[decision.id];scope=decision.applies_to
            scope_text=('対象: '+', '.join([*scope.resources,*scope.types])+('\n環境: '+', '.join(scope.environments) if scope.environments else '')) if scope else '適用範囲は未宣言'
            add('設計方針',[decision.id,decision.section+'\n'+titles[decision.section],decision.title+'\n'+decision.status,decision.statement,
                decision.rationale+('\n代替案: '+' / '.join(decision.alternatives) if decision.alternatives else ''),
                '要件: '+(', '.join(decision.requirement_ids) or 'なし')+'\n検査: '+(', '.join(decision.policy_ids) or 'なし')+('\n結果: '+text(row['policy_results']) if row['policy_results'] else ''),
                scope_text+'\n具体化: '+row['realization']+' / '+row['mechanization'],', '.join(decision.sources)])
        for link in trace['links']:
            add('方針トレース',[link['decision_id'],link['item_id']+'\n'+link['path'],link['status']+' / '+(link['mode'] or '-'),link['value_resolution'],link['input_pointer']])
        for record in trace.get('expansions',[]):
            add('方針展開',[record['expansion_id'],record['decision_id'],record['item_id'],record['result'],text(record['setting']),record['input_pointer']])
        for f in trace['findings']:
            add('方針指摘',[f['finding_id'],f['kind']+' / '+f['code'],'\n'.join(f[k] for k in ('decision_id','section_id','requirement_id') if f.get(k)),
                '\n'.join(f[k] for k in ('item_id','path','resource_id') if f.get(k)),f['message']])
    if document.model_version=='1.1':
        add('概要',['@comparison-legend','比較と値の解決','DECIDEDは入力方針の決定。RESOLVEDは記載値が具体値であること。MATCHは記載された具体値の一致のみ。EXPRESSION_MATCHは式の一致のみで実値は未解決。REFERENCE_MATCHは参照宣言の一致のみ。いずれも設計合格や接続・デプロイの成立を意味しません。'])
        add('概要',['@catalog','CFNスキーマ',document.catalog_provenance])
        add('概要',['@volume','この巻のサービス設定',', '.join(selected) or '共通索引（設定はサービス分冊）'])
        for path in document.read_only_omissions:add('差異',['@read-only','READ_ONLY','','',path])
    for ws in book:
        ws.auto_filter.ref=ws.dimensions;ws.print_area=ws.dimensions
    return book
