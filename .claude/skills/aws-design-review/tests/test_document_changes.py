"""Workbook proposals become reviewed input changes only through propose/apply."""
import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from aws_design_sheet.document_changes import apply, main, propose
from aws_design_sheet.document_cli import generate
from aws_design_sheet.document_v11 import SchemaInput

EXAMPLE=Path(__file__).resolve().parents[1]/'examples/basic-design/input.json'
VALUE,REASON,SCOPE='修正案（値）','修正理由','適用先（共通/環境）'


def fill(path,sheet,item_id,values,row=1):
    """Write edit columns on the n-th row of an item, as a reviewer would."""
    book=load_workbook(path);ws=book[sheet];header=[c.value for c in ws[1]]
    rows=[r for r in ws.iter_rows(min_row=2) if r[0].value==item_id]
    for name,value in values.items():ws.cell(rows[row-1][0].row,header.index(name)+1,value)
    book.save(path)


@pytest.fixture
def generated(tmp_path):
    source=tmp_path/'input.json';source.write_bytes(EXAMPLE.read_bytes())
    generate(source,'dev',tmp_path/'gen')
    return source,tmp_path/'gen',tmp_path/'gen/services-01.xlsx'


def test_round_trip_records_approval_and_regenerates(generated,tmp_path):
    source,folder,book=generated;before=hashlib.sha256(source.read_bytes()).hexdigest()
    fill(book,'Logs設定','app-logs.RetentionInDays',{VALUE:'60',REASON:'障害調査に2か月必要',SCOPE:'環境'})
    fill(book,'EC2設定','subnet-c.CidrBlock',{VALUE:'10.0.3.0/24',REASON:'AZ-cの割当確定'})
    proposal=propose(source,[folder/'design.xlsx',book])
    assert proposal['summary']=={'changes':2,'rejected':0,'rejected_codes':{}}
    by_item={c['item_id']:c for c in proposal['changes']}
    assert by_item['app-logs.RetentionInDays']['scope']=='ENVIRONMENT' and by_item['app-logs.RetentionInDays']['before'] is None
    assert by_item['subnet-c.CidrBlock']['scope']=='COMMON'
    assert by_item['subnet-c.CidrBlock']['after']=={'mode':'LITERAL','value':'10.0.3.0/24','decision_ids':['dec-vpc-cidr']}
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before
    path=tmp_path/'proposal.json';path.write_text(json.dumps(proposal,ensure_ascii=False),encoding='utf-8')
    out=tmp_path/'input-v2.json'
    assert main(['apply',str(source),str(path),'--output',str(out),'--approver','設計責任者','--all','--generate',str(tmp_path/'gen2')])==0
    data=json.loads(out.read_text(encoding='utf-8'));SchemaInput.model_validate(data)
    assert {c['change_id'] for c in data['change_log']}=={c['change_id'] for c in proposal['changes']}
    assert all(c['approved_by']=='設計責任者' and c['proposal_sha256'] for c in data['change_log'])
    design=json.loads((tmp_path/'gen2/design.json').read_text(encoding='utf-8'))
    items={i['item_id']:i for i in design['items']}
    assert items['app-logs.RetentionInDays']['origin']=='environment' and items['app-logs.RetentionInDays']['expected']['value']==60
    assert items['subnet-c.CidrBlock']['expected']['value']=='10.0.3.0/24'
    trace=design['decision_trace']
    assert not any(f['code']=='REALIZING_SETTING_UNDECIDED' for f in trace['findings'])
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before


def test_selected_change_only_and_undecided_word(generated,tmp_path):
    source,folder,book=generated
    fill(book,'Logs設定','app-logs.RetentionInDays',{VALUE:'未決定',REASON:'保持期間を再検討'})
    fill(book,'EC2設定','vpc.EnableDnsSupport',{VALUE:False,REASON:'名前解決は別経路'})
    proposal=propose(source,[book]);path=tmp_path/'p.json';path.write_text(json.dumps(proposal),encoding='utf-8')
    undecided=next(c for c in proposal['changes'] if c['item_id']=='app-logs.RetentionInDays')
    assert undecided['after']=={'status':'UNDECIDED','decision_ids':['dec-log-retention']}
    result=apply(source,path,tmp_path/'v2.json','reviewer',[undecided['change_id']])
    assert result['not_applied']==[c['change_id'] for c in proposal['changes'] if c is not undecided]
    data=json.loads((tmp_path/'v2.json').read_text(encoding='utf-8'))
    assert next(r for r in data['resources'] if r['id']=='app-logs')['settings']['RetentionInDays']['status']=='UNDECIDED'
    assert next(r for r in data['resources'] if r['id']=='vpc')['settings']['EnableDnsSupport']['value'] is True


@pytest.mark.parametrize('item,sheet,values,code',[
    ('app-logs.LogGroupName','Logs設定',{VALUE:'=CONCAT("a")',REASON:'x'},'FORMULA'),
    ('app-logs.LogGroupName','Logs設定',{VALUE:'/portal/new'},'REASON_REQUIRED'),
    ('app-logs.LogGroupName','Logs設定',{REASON:'値なし'},'VALUE_REQUIRED'),
    # Fictional key ID, assembled so repository secret scanners do not flag the fixture.
    ('app-logs.LogGroupName','Logs設定',{VALUE:'AK'+'IA'+'ABCDEFGHIJKLMNOP',REASON:'x'},'SECRET_LIKE'),
    ('app-logs.LogGroupName','Logs設定',{VALUE:'/x',REASON:'x',SCOPE:'全体'},'UNKNOWN_SCOPE'),
    ('assets.BucketEncryption','S3設定',{VALUE:'{}',REASON:'x'},'UNSUPPORTED_TARGET_MODE'),
    ('app-logs.RetentionInDays','Logs設定',{VALUE:'"abc"',REASON:'x'},'INVALID_FOR_SCHEMA'),
    ('vpc.CidrBlock','EC2設定',{VALUE:'10.0.0.0/16',REASON:'同値'},'UNCHANGED'),
])
def test_unsafe_or_invalid_proposals_are_rejected(generated,item,sheet,values,code):
    source,_,book=generated;fill(book,sheet,item,values)
    proposal=propose(source,[book])
    assert proposal['changes']==[] and [r['code'] for r in proposal['rejected']]==[code]
    assert 'AKIA' not in json.dumps(proposal)


def test_ambiguous_cells_and_stale_inputs_fail(generated,tmp_path,capsys):
    source,folder,book=generated
    fill(book,'S3設定','assets.BucketEncryption',{VALUE:'{}',REASON:'x'})
    fill(book,'S3設定','assets.BucketEncryption',{VALUE:'[]'},row=2)
    assert [r['code'] for r in propose(source,[book])['rejected']]==['AMBIGUOUS_CELLS']
    proposal=tmp_path/'p.json'
    source.write_bytes(source.read_bytes()+b'\n')
    assert main(['propose',str(source),str(folder),'--output',str(proposal)])==2
    assert 'STALE_WORKBOOK' in capsys.readouterr().out and not proposal.exists()


def test_apply_guards(generated,tmp_path,capsys):
    source,_,book=generated
    fill(book,'Logs設定','app-logs.LogGroupName',{VALUE:'/portal/app2',REASON:'命名規則の変更'})
    path=tmp_path/'p.json';path.write_text(json.dumps(propose(source,[book])),encoding='utf-8')
    with pytest.raises(ValueError,match='approver'):apply(source,path,tmp_path/'a.json',' ')
    with pytest.raises(ValueError,match='unknown or duplicate'):apply(source,path,tmp_path/'a.json','r',['chg-0000000000000000'])
    with pytest.raises(ValueError,match='differ'):apply(source,path,source,'r')
    assert main(['apply',str(source),str(path),'--output',str(tmp_path/'a.json'),'--approver','r'])==2
    assert not (tmp_path/'a.json').exists()
    apply(source,path,tmp_path/'a.json','r')
    with pytest.raises(ValueError,match='already exists'):apply(source,path,tmp_path/'a.json','r')
    data=json.loads((tmp_path/'a.json').read_text(encoding='utf-8'))
    data['change_log'].append(data['change_log'][0])
    with pytest.raises(ValidationError):SchemaInput.model_validate(data)
