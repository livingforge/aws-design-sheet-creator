"""Bridge controls: equality never implies AWS validity or complete coverage."""
import copy
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from aws_design_sheet.document_cli import generate, main
from aws_design_sheet.document_evaluation import adapt, evaluate
from aws_design_sheet.document_normalize import normalize


def lit(value): return {'mode': 'LITERAL', 'value': value}


def spec(kind='AWS::Lambda::Function', settings=None):
    return {'schema_version': '1.1', 'system': {'id': 'demo', 'name': 'Bridge'},
            'environments': [{'id': 'dev', 'name': 'Development', 'region': 'ap-northeast-1'},
                             {'id': 'prod', 'name': 'Production', 'region': 'ap-northeast-1'}],
            'sources': [{'id': 'source', 'title': 'Design decision', 'locator': 'spec:1'}],
            'requirements': [{'id': 'req', 'title': 'Business intention', 'text': 'Not an executable rule', 'resources': ['resource'], 'sources': ['source']}],
            'resources': [{'id': 'resource', 'name': 'Resource', 'type': kind, 'cfn_logical_id': 'Resource',
                           'sources': ['source'], 'settings': settings if settings is not None else {
                               'Runtime': {**lit('nodejs22.x'), 'sources': ['source']},
                               'Handler': lit('index.handler'), 'Role': lit('arn:aws:iam::123456789012:role/Test'),
                               'Code': lit({'ZipFile': 'exports.handler=async()=>{};'})}}]}


def norm(data, cfn=None, env='dev'):
    return normalize(json.dumps(data).encode(), env, json.dumps(cfn).encode() if cfn else None)


def cfn_for(data):
    resource = data['resources'][0]
    return {'Resources': {'Resource': {'Type': resource['type'], 'Properties': {
        k: v['value'] for k, v in resource['settings'].items() if v.get('mode') == 'LITERAL'}}}}


def findings(result, rule): return [r for r in result.results if r['rule_id'] == rule and r['verdict'] == 'FAIL']


@pytest.mark.parametrize('kind,settings,field,invalid,valid,rule', [
    ('AWS::Lambda::Function', None, 'Runtime', 'invalid-runtime', 'nodejs22.x', 'CFN_LINT.E3030'),
    ('AWS::ApiGateway::DomainName', {'DomainName': lit('api.example.com'), 'SecurityPolicy': lit('TLS_1_2')},
     'SecurityPolicy', 'TLS_INVALID', 'TLS_1_2', 'CFN_LINT.E9001'),
])
@pytest.mark.parametrize('with_cfn', [False, True])
def test_same_invalid_value_and_repair(kind, settings, field, invalid, valid, rule, with_cfn):
    data = spec(kind, settings)
    data['resources'][0]['settings'][field] = {**lit(invalid), 'sources': ['source']}
    doc = norm(data, cfn_for(data) if with_cfn else None)
    item = next(i for i in doc.items if i.name == field)
    if with_cfn: assert item.comparison == 'MATCH'
    result = evaluate(doc)
    assert result.status == 'COMPLETE'
    hits = findings(result, rule)
    assert hits and hits[0]['trace'][0]['item_id'] == 'resource.' + field
    assert hits[0]['trace'][0]['source_ids'] == ['source']
    assert hits[0]['requirement_ids'] == ['req']
    assert result.assessment == 'FAIL'
    data['resources'][0]['settings'][field] = lit(valid)
    repaired = evaluate(norm(data, cfn_for(data) if with_cfn else None))
    assert repaired.status == 'COMPLETE' and not findings(repaired, rule)


def test_required_absent_on_both_sides_and_no_cfn_borrowing():
    data = spec(settings={})
    doc = norm(data, cfn_for(data)); result = evaluate(doc)
    assert not doc.items and doc.checklist.status == 'INCOMPLETE'
    assert findings(result, 'CFN_LINT.E3003')
    supplied = cfn_for(spec())
    doc = norm(data, supplied)
    assert all(i.origin == 'cfn_only' for i in doc.items)
    assert findings(evaluate(doc), 'CFN_LINT.E3003')
    assert not adapt(doc)[0].resources[0].fields


def test_unknowns_composites_references_and_secret_never_resolve():
    data = spec()
    s = data['resources'][0]['settings']
    s['Role'] = {'mode': 'EXTERNAL_REF', 'ref': 'external://role'}
    s['Runtime'] = {'mode': 'EXPRESSION', 'expression': {'Ref': 'RuntimeParameter'}}
    s['MemorySize'] = {'mode': 'DEFAULT'}
    s['Timeout'] = {'status': 'UNDECIDED'}
    s['Environment'] = {'mode': 'COMPOSITE', 'value': {'Variables': {
        'PUBLIC': 'hello', 'Password': {'$design': {'mode': 'SECRET_REF', 'classification': 'SECRET', 'ref': 'secret://probe'}}}}}
    doc = norm(data); design, mappings, coverage = adapt(doc)
    assert design.resources[0].field('/properties/Runtime').state.value == 'UNRESOLVED'
    assert next(i for i in doc.items if i.name == 'Runtime').expected.expression == {'Ref': 'RuntimeParameter'}
    assert any(m['path'] == '/properties/Environment/Variables/Password' and m['mode'] == 'SECRET_REF' for m in mappings)
    result = evaluate(doc)
    assert result.status == 'COMPLETE' and result.assessment != 'NO_FINDINGS'
    for template in result.provenance['synthetic_templates']:
        assert template['synthetic'] and not template['deployable']
        assert all('Default' not in p for p in template['template'].get('Parameters', {}).values())
    assert 'RuntimeParameter' not in json.dumps(result.provenance['synthetic_templates'])


def test_internal_reference_is_preserved_without_stack_or_attribute_invention():
    data = spec('AWS::EC2::VPC', {'CidrBlock': lit('10.0.0.0/16')})
    data['schema_version'] = '1.0'
    data['resources'].append({'id': 'subnet', 'name': 'Subnet', 'type': 'AWS::EC2::Subnet',
                              'settings': {'VpcId': {'mode': 'RESOURCE_REF', 'ref': 'resource'}, 'CidrBlock': lit('10.0.1.0/24')}})
    design, _, _ = adapt(norm(data))
    assert design.relations[0].target_resource_id == 'resource'
    assert all(r.template is None for r in design.resources)
    result = evaluate(norm(data)); assert result.status == 'COMPLETE'
    assert any(c['kind'] == 'CFN_LINT_CROSS_TEMPLATE' for c in result.coverage)
    assert len(result.provenance['synthetic_templates']) == 2


def test_environment_override_ids_and_sources():
    data = spec()
    data['environments'][1]['overrides'] = [{'resource_id': 'resource', 'settings': {'Runtime': {**lit('invalid-runtime'), 'sources': ['source']}}}]
    assert not findings(evaluate(norm(data)), 'CFN_LINT.E3030')
    result = evaluate(norm(data, env='prod'))
    hit = findings(result, 'CFN_LINT.E3030')[0]
    assert hit['environment_id'] == 'prod'
    assert hit['trace'][0]['input_pointer'].startswith('/environments/1/overrides/0/')
    assert hit['trace'][0]['row_id'] == 'prod/resource.Runtime'


@pytest.mark.parametrize('mode', ['exception', 'failed_status'])
@pytest.mark.parametrize('stage', ['run_checker', 'run_linter'])
def test_evaluator_failure_never_passes(monkeypatch, stage, mode):
    def fail(_):
        if mode == 'exception': raise RuntimeError('SECRET_CANARY_EXCEPTION')
        return {'status': 'FAILED'}
    monkeypatch.setattr('aws_design_sheet.document_evaluation.' + stage, fail)
    result = evaluate(norm(spec()))
    assert result.status == 'FAILED' and result.assessment == 'ERROR'
    assert 'SECRET_CANARY_EXCEPTION' not in result.model_dump_json()


def test_unsupported_region_is_skipped():
    data = spec('AWS::EC2::VPC', {'CidrBlock': lit('10.0.0.0/16')})
    data['schema_version'] = '1.0'; data['environments'][0]['region'] = 'us-west-2'
    result = evaluate(norm(data))
    assert result.status == 'SKIPPED' and result.assessment == 'NOT_EVALUATED'
    assert not result.results


def test_type_error_is_structure_error_and_condition_rules_are_independent(tmp_path, capsys):
    data = spec(); data['resources'][0]['settings']['Timeout'] = lit('wrong-type')
    path = tmp_path / 'bad.json'; path.write_text(json.dumps(data))
    assert main([str(path), '--environment', 'dev', '--output', str(tmp_path / 'out'), '--evaluate']) == 2
    assert json.loads(capsys.readouterr().out)['generation_status'] == 'STRUCTURE_ERROR'
    data = spec(); data['resources'][0]['settings']['PackageType'] = lit('Image')
    result = evaluate(norm(data))
    assert any(r['verdict'] == 'FAIL' and r['rule_id'].startswith('CFN_LINT.') for r in result.results)


def read_evaluation_book(path):
    book = load_workbook(path, data_only=False)
    entries = {}
    for row in book['AWS評価根拠'].iter_rows(min_row=2, values_only=True):
        entries.setdefault(row[0], '')
        entries[row[0]] += row[2] or ''
    decoded = {key: json.loads(value) for key, value in entries.items()}
    metadata = decoded.pop('@evaluation')
    metadata['results'] = [r for r in decoded.values() if 'verdict' in r]
    metadata['coverage'] = [r for r in decoded.values() if 'verdict' not in r]
    for name in ('AWS評価', 'AWS未検査'):
        for row in book[name].iter_rows(min_row=2):
            assert row[0].hyperlink and row[0].value in decoded
            location = int(row[0].hyperlink.target.rsplit('A', 1)[1])
            assert book['AWS評価根拠'].cell(location, 1).value == row[0].value
    assert all(c.data_type != 'f' for sheet in book for row in sheet for c in row)
    book.close()
    return metadata


def test_json_xlsx_roundtrip_secrets_and_cli_failure(tmp_path, monkeypatch):
    data = spec(); path = tmp_path / 'input.json'; path.write_text(json.dumps(data))
    cfn = cfn_for(data); cfn['Resources']['Resource']['Properties']['Environment'] = {'Variables': {'Password': 'SECRET_CANARY_VALUE'}}
    cpath = tmp_path / 'cfn.json'; cpath.write_text(json.dumps(cfn))
    output = tmp_path / 'evaluated'; summary = generate(path, 'dev', output, cpath, evaluate_design=True)
    result = json.loads((output / 'evaluation.json').read_text(encoding='utf-8'))
    assert summary['generation_status'] == 'GENERATED' and summary['evaluation_status'] == 'COMPLETE'
    assert read_evaluation_book(output / 'design.xlsx') == result
    for p in output.iterdir():
        if p.suffix == '.json': assert 'SECRET_CANARY_VALUE' not in p.read_text(encoding='utf-8')
        if p.suffix == '.xlsx':
            with ZipFile(p) as z:
                assert all(b'SECRET_CANARY_VALUE' not in z.read(n) for n in z.namelist())
    monkeypatch.setattr('aws_design_sheet.document_evaluation.run_checker', lambda _: {'status': 'FAILED'})
    assert main([str(path), '--environment', 'dev', '--output', str(tmp_path / 'failed'), '--evaluate']) == 3
    assert json.loads((tmp_path / 'failed/summary.json').read_text())['aws_assessment'] == 'ERROR'


def test_default_generation_does_not_claim_evaluation(tmp_path):
    path = tmp_path / 'input.json'; path.write_text(json.dumps(spec()))
    summary = generate(path, 'dev', tmp_path / 'plain')
    assert summary['evaluation_status'] == 'NOT_EVALUATED'
    assert not (tmp_path / 'plain/evaluation.json').exists()
