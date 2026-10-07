"""Retrospective evaluation of past designs against their known issues."""
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.evaluate_retrospective import LABELS, evaluate_case, load_case


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "examples/retrospective/sample-case"


def run(directory, update_triage=True):
    return evaluate_case(directory, schema_dir=ROOT / "schemas",
                         profile=ROOT / "profiles/vpc-subnet.json", update_triage=update_triage)


def test_sample_case_outcomes_and_scope():
    report = run(SAMPLE, update_triage=False)
    outcomes = {issue["id"]: issue["outcome"] for issue in report["issues"]}
    assert outcomes == {"K1": "DETECTED", "K2": "DETECTED", "K3": "DETECTED", "K4": "FLAGGED",
                        "K5": "MISSED", "K6": "MISSED", "K7": "OUT_OF_SCOPE"}
    assert report["fictional"] is True
    assert report["known_issues"] == {"total": 7, "in_scope": 6, "detected": 3, "flagged": 1,
                                      "missed": 2, "out_of_scope": 1}
    assert report["recall_in_scope_any"] == 0.667
    assert report["triage"]["fail_precision"] == 1.0


def test_triage_labels_survive_reruns(tmp_path):
    case = tmp_path / "case"
    shutil.copytree(SAMPLE, case)
    first = run(case)
    assert first["triage"]["labeled"] == 0 and first["triage"]["pending"] > 0
    workbook = load_workbook(case / "triage.xlsx")
    sheet = workbook["未対応の指摘"]
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=11, value=LABELS[2])
        sheet.cell(row=row, column=12, value="note")
    workbook.save(case / "triage.xlsx")
    second = run(case)
    assert second["triage"]["pending"] == 0
    assert second["triage"]["useful_rate"] == 0.0
    rows = list(load_workbook(case / "triage.xlsx")["未対応の指摘"].iter_rows(min_row=2, values_only=True))
    assert rows and all(row[10] == LABELS[2] and row[11] == "note" for row in rows)


def test_case_definition_is_validated(tmp_path):
    case = tmp_path / "case"
    shutil.copytree(SAMPLE, case)
    data = json.loads((case / "case.json").read_text(encoding="utf-8"))
    data["known_issues"][0]["targets"][0].pop("path")
    (case / "case.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValueError, match="needs a path"):
        load_case(case)
    data["known_issues"][0]["targets"] = []
    (case / "case.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValueError, match="needs targets"):
        load_case(case)
