"""Retrospective evaluation on past project designs with known issues.

Each case directory holds ``case.json`` (scope, inputs, conversion effort and
the issues the project actually found) and the design text under ``inputs/``.
Findings are matched to known issues by resource and property path. Findings
that match no known issue go to ``triage.xlsx`` in the case directory, where a
reviewer labels them; labels already entered are read back on the next run, so
usefulness can be measured without losing earlier judgments.
"""
from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from aws_design_sheet.excel_report import put, table, text
from aws_design_sheet.extractor import TextSource
from aws_design_sheet.runner import run_text


ROOT = Path(__file__).resolve().parents[1]
FINDING_VERDICTS = ("FAIL", "NEEDS_REVIEW", "ERROR")
LABELS = ("有用：新たな問題", "有用：確認として妥当", "不要：誤検知・ノイズ", "保留")
TRIAGE_HEADERS = ["キー", "判定", "重大度", "ルール", "種別", "リソース", "項目", "理由", "実際の値",
                  "根拠行", "レビュー判定", "メモ"]
FOUND_BY = ("design_review", "deploy", "test", "operation", "incident")


def load_case(directory: Path) -> dict:
    case = json.loads((directory / "case.json").read_text(encoding="utf-8"))
    for key in ("case_id", "project", "environment", "account", "inputs", "known_issues"):
        if key not in case:
            raise ValueError(f"{directory}: case.json lacks {key}")
    ids = [issue["id"] for issue in case["known_issues"]]
    if len(ids) != len(set(ids)):
        raise ValueError(f"{directory}: duplicate known issue ID")
    for issue in case["known_issues"]:
        if "in_scope" not in issue:
            raise ValueError(f"{directory}: issue {issue['id']} needs in_scope")
        if issue["in_scope"] and not issue.get("targets"):
            raise ValueError(f"{directory}: in-scope issue {issue['id']} needs targets")
        if any(not target.get("path") for target in issue.get("targets", [])):
            raise ValueError(f"{directory}: issue {issue['id']} target needs a path ('*' for the whole resource)")
        if issue.get("found_by") not in FOUND_BY:
            raise ValueError(f"{directory}: issue {issue['id']} found_by must be one of {FOUND_BY}")
    return case


def finding_key(item: dict, resource) -> str:
    where = f"{resource.type}/{resource.name}" if resource else "-"
    return "|".join((item["rule_id"], where, item.get("path") or "", item["verdict"]))


def overlaps(finding: str | None, target: str) -> bool:
    if target == "*":
        return True
    if not finding:
        return False
    return finding == target or finding.startswith(target + "/") or target.startswith(finding + "/")


def matches(item: dict, resource, target: dict) -> bool:
    return (resource is not None and resource.type == target["type"]
            and resource.name == target["name"] and overlaps(item.get("path"), target.get("path")))


def read_labels(path: Path) -> dict[str, tuple[str, str]]:
    if not path.exists():
        return {}
    sheet = load_workbook(path, read_only=True)["未対応の指摘"]
    rows = sheet.iter_rows(min_row=2, values_only=True)
    return {row[0]: (row[10] or "", row[11] or "") for row in rows if row and row[0]}


def write_triage(path: Path, rows: list[list]):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "未対応の指摘"
    table(sheet, TRIAGE_HEADERS, rows, [10, 14, 10, 34, 28, 22, 30, 50, 30, 16, 22, 30],
          wrapped={8, 12})
    sheet.column_dimensions["A"].hidden = True
    choices = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    sheet.add_data_validation(choices)
    choices.add(f"K2:K{max(len(rows) + 1, 2)}")
    guide = workbook.create_sheet("判定の基準")
    for row, line in enumerate((
            "既知の問題に当たらなかった指摘を 1 件ずつ判定する。",
            "有用：新たな問題 … 当時見つかっていなかった設計上の問題。",
            "有用：確認として妥当 … 問題ではないが、レビューで確認すべき事項として妥当。",
            "不要：誤検知・ノイズ … 設計に問題がなく、確認の必要もない。",
            "保留 … 判断材料が足りない。メモに理由を書く。"), 1):
        put(guide, row, 1, line)
    guide.column_dimensions["A"].width = 90
    workbook.save(path)


def evaluate_case(directory: Path, *, schema_dir: Path, profile: Path, update_triage: bool) -> dict:
    case = load_case(directory)
    sources = [TextSource(id=f"doc-{index}", name=name, version="1",
                          text=(directory / name).read_bytes().decode("utf-8-sig"))
               for index, name in enumerate(case["inputs"], 1)]
    design, result = run_text(sources, project=case["project"], environment=case["environment"],
                              account=case["account"], region=case.get("region", "ap-northeast-1"),
                              schema_dir=schema_dir, profile_path=profile)
    by_id = {resource.id: resource for resource in design.resources}
    evidence = {item.id: item for item in design.evidence}
    documents = {item.id: item.name for item in design.documents}
    findings = [item for item in result["results"] if item["verdict"] in FINDING_VERDICTS]

    issues, matched = [], set()
    for issue in case["known_issues"]:
        if not issue["in_scope"]:
            # Out-of-scope issues cannot be found in the design text; a finding on the
            # same resource would be a coincidence, so they are not matched.
            issues.append({"id": issue["id"], "summary": issue.get("summary", ""),
                           "found_by": issue["found_by"], "in_scope": False,
                           "outcome": "OUT_OF_SCOPE", "findings": []})
            continue
        hits = [(index, item) for index, item in enumerate(findings)
                if any(matches(item, by_id.get(item["resource_id"]), target) for target in issue["targets"])]
        matched.update(index for index, _ in hits)
        verdicts = {item["verdict"] for _, item in hits}
        outcome = "DETECTED" if "FAIL" in verdicts else "FLAGGED" if hits else "MISSED"
        issues.append({"id": issue["id"], "summary": issue.get("summary", ""),
                       "found_by": issue["found_by"], "in_scope": issue["in_scope"], "outcome": outcome,
                       "findings": sorted({finding_key(item, by_id.get(item["resource_id"]))
                                           for _, item in hits})})

    labels = read_labels(directory / "triage.xlsx")
    unmatched, rows = [], []
    for index, item in enumerate(findings):
        if index in matched:
            continue
        resource = by_id.get(item["resource_id"])
        key = finding_key(item, resource)
        label, note = labels.get(key, ("", ""))
        lines = sorted({f"{documents[evidence[e].document_id]}:{evidence[e].start_line}"
                        for e in item.get("evidence_ids", []) if e in evidence})
        unmatched.append({"key": key, "verdict": item["verdict"], "rule_id": item["rule_id"],
                          "label": label, "note": note})
        rows.append([key, item["verdict"], item.get("severity", "ERROR"), item["rule_id"],
                     resource.type if resource else "", resource.name if resource else "",
                     item.get("path") or "", item.get("reason") or "", text(item.get("actual")),
                     ", ".join(lines), label, note])
    if update_triage:
        write_triage(directory / "triage.xlsx", rows)

    in_scope = [issue for issue in issues if issue["in_scope"]]
    outcome_counts = Counter(issue["outcome"] for issue in in_scope)
    fails = [item for item in findings if item["verdict"] == "FAIL"]
    matched_fails = sum(1 for index, item in enumerate(findings) if index in matched and item["verdict"] == "FAIL")
    labeled = [item for item in unmatched if item["label"] in LABELS[:3]]
    labeled_fails = [item for item in labeled if item["verdict"] == "FAIL"]
    useful_fails = matched_fails + sum(item["label"].startswith("有用") for item in labeled_fails)
    reviewed_fails = matched_fails + len(labeled_fails)
    useful = sum(item["label"].startswith("有用") for item in labeled)
    unprocessed = sum(item["kind"] == "UNPROCESSED_TEXT" for item in result["coverage"])
    resources = len(design.resources)
    return {
        "case_id": case["case_id"], "fictional": bool(case.get("fictional")),
        "status": result["status"], "resources": resources, "unprocessed_lines": unprocessed,
        "conversion": case.get("conversion", {}),
        "verdicts": result["summary"],
        "known_issues": {"total": len(issues), "in_scope": len(in_scope),
                         **{key.lower(): outcome_counts.get(key, 0) for key in ("DETECTED", "FLAGGED", "MISSED")},
                         "out_of_scope": len(issues) - len(in_scope)},
        "recall_in_scope_fail": round(outcome_counts["DETECTED"] / len(in_scope), 3) if in_scope else None,
        "recall_in_scope_any": round((outcome_counts["DETECTED"] + outcome_counts["FLAGGED"]) / len(in_scope), 3)
        if in_scope else None,
        "unmatched_findings": dict(Counter(item["verdict"] for item in unmatched)),
        "triage": {"labeled": len(labeled), "pending": len(unmatched) - len(labeled),
                   "labels": dict(Counter(item["label"] for item in unmatched if item["label"])),
                   "fail_precision": round(useful_fails / reviewed_fails, 3) if reviewed_fails else None,
                   "fail_reviewed": reviewed_fails, "fail_total": len(fails),
                   "useful_rate": round(useful / len(labeled), 3) if labeled else None},
        "needs_review_per_100_resources": round(100 * result["summary"].get("NEEDS_REVIEW", 0) / resources, 1)
        if resources else None,
        "issues": issues}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("cases", nargs="+", type=Path, help="case directories with case.json")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--schemas", type=Path, default=ROOT / "schemas")
    parser.add_argument("--profile", type=Path, default=ROOT / "profiles/vpc-subnet.json")
    parser.add_argument("--no-triage-update", action="store_true",
                        help="do not rewrite triage.xlsx in the case directories")
    args = parser.parse_args()
    reports = [evaluate_case(case, schema_dir=args.schemas, profile=args.profile,
                             update_triage=not args.no_triage_update) for case in args.cases]
    real = [report for report in reports if not report["fictional"]]
    totals = Counter()
    for report in real:
        for key in ("in_scope", "detected", "flagged", "missed"):
            totals[key] += report["known_issues"][key]
    summary = {"cases": len(reports), "real_cases": len(real), **totals,
               "recall_in_scope_any": round((totals["detected"] + totals["flagged"]) / totals["in_scope"], 3)
               if totals["in_scope"] else None}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps({"summary": summary, "cases": reports}, ensure_ascii=False, indent=2) + "\n",
                           encoding="utf-8")
    for report in reports:
        print(json.dumps({key: report[key] for key in ("case_id", "known_issues", "recall_in_scope_any",
                                                       "unmatched_findings", "triage")}, ensure_ascii=False))
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
