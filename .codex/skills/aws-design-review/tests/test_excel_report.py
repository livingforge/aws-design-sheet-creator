import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from aws_design_sheet.excel_report import build_workbook, main
from aws_design_sheet.extractor import TextSource
from aws_design_sheet.runner import run_text

ROOT = Path(__file__).resolve().parents[1]


def run(text):
    return run_text([TextSource(id="doc-1", name="design.txt", version="1", text=text)],
                    project="p", environment="prod", account="111111111111",
                    region="ap-northeast-1", schema_dir=ROOT / "schemas",
                    profile_path=ROOT / "profiles/vpc-subnet.json")


def rows(ws):
    header, *body = ws.iter_rows(values_only=True)
    return [dict(zip(header, row)) for row in body]


def test_results_sheet_joins_design_and_orders_failures_first():
    design, result = run((ROOT / "examples/review-demo.txt").read_text(encoding="utf-8"))
    wb = build_workbook(result, design)
    assert wb.sheetnames == ["概要", "検査結果", "未検査範囲", "設計値"]
    findings = rows(wb["検査結果"])
    assert len(findings) == len(result["results"])
    assert [f["判定"] for f in findings[:3]] == ["FAIL", "FAIL", "NEEDS_REVIEW"]
    outside = next(f for f in findings if f["ルール ID"] == "CIDR_CONTAINMENT"
                   and f["判定"] == "FAIL")
    assert (outside["リソース種別"], outside["リソース名"]) == ("AWS::EC2::Subnet", "private-c")
    assert outside["実際の値"] == "10.1.2.0/24"
    assert outside["根拠"].startswith("design.txt L3: ")
    assert wb["検査結果"]["B2"].fill.fgColor.rgb.endswith("FFC7CE")
    assert len(rows(wb["未検査範囲"])) == len(result["coverage"])
    design_rows = rows(wb["設計値"])
    reference = next(r for r in design_rows if r["リソース名"] == "primary"
                     and r["項目パス"] == "/properties/DBSubnetGroupName")
    assert (reference["状態"], reference["値"]) == ("REFERENCE", "→ AWS::RDS::DBSubnetGroup/db-subnets")


def test_design_sheet_keeps_false_and_blocks_formulas():
    design, result = run('VPC main: CidrBlock=10.0.0.0/16; EnableDnsSupport=false\n'
                         'AWS::S3::Bucket logs: BucketName="=cmd|x"\n')
    ws = build_workbook(result, design)["設計値"]
    values = {(r["リソース名"], r["項目パス"]): r for r in rows(ws)}
    assert values[("main", "/properties/EnableDnsSupport")]["値"] == "false"
    assert values[("main", "/properties/EnableDnsSupport")]["状態"] == "KNOWN"
    cell = next(c for c in ws["F"] if c.value == "=cmd|x")
    assert cell.data_type == "s"


def test_rejects_result_from_other_design():
    design, _ = run("VPC main: CidrBlock=10.0.0.0/16\n")
    _, other = run("VPC main: CidrBlock=10.1.0.0/16\n")
    with pytest.raises(ValueError, match="does not match"):
        build_workbook(other, design)


def test_failed_result_without_design(tmp_path):
    result_path = tmp_path / "result.json"
    result_path.write_text(json.dumps({"status": "FAILED", "diagnostic": "boom",
                                       "results": [], "coverage": []}), encoding="utf-8")
    assert main([str(result_path), "--output", str(tmp_path / "out.xlsx")]) == 0
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb.sheetnames == ["概要", "検査結果", "未検査範囲"]
    summary = {row[0]: row[1] for row in wb["概要"].iter_rows(values_only=True) if row[0]}
    assert summary["処理状態"] == "FAILED"
    assert summary["診断情報"] == "boom"
