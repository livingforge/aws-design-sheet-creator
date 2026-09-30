"""Write check results (and optionally the checked design data) to an Excel workbook."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Design

VERDICT_ORDER = ["ERROR", "FAIL", "NEEDS_REVIEW", "NOT_APPLICABLE", "PASS"]
VERDICT_STYLES = {  # (fill, font) colours
    "PASS": ("C6EFCE", "006100"),
    "FAIL": ("FFC7CE", "9C0006"),
    "NEEDS_REVIEW": ("FFEB9C", "9C5700"),
    "NOT_APPLICABLE": ("EDEDED", "595959"),
    "ERROR": ("E4DFEC", "5B2C6F"),
}
STATE_STYLES = {
    "KNOWN": VERDICT_STYLES["PASS"],
    "MISSING": VERDICT_STYLES["FAIL"],
    "CONFLICT": VERDICT_STYLES["FAIL"],
    "INFERRED": VERDICT_STYLES["NEEDS_REVIEW"],
    "UNRESOLVED": VERDICT_STYLES["NEEDS_REVIEW"],
    "NOT_APPLICABLE": VERDICT_STYLES["NOT_APPLICABLE"],
    "REFERENCE": ("DDEBF7", "1F4E78"),  # logical reference to another resource
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11, color="1F4E78")
TOP = Alignment(vertical="top")
WRAP = Alignment(wrap_text=True, vertical="top")
MAX_CELL = 32767
NOTE = "検査済み範囲の合格は、設計全体の妥当性や AWS へのデプロイ可否を保証しない。"


def text(value: Any) -> str:
    """Render a JSON value for a cell; None stays empty and false/0 stay visible."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False)


def put(ws, row: int, column: int, value: Any, *, style: tuple[str, str] | None = None,
        wrap: bool = False):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL]
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # design text is data, never a formula
    if style:
        cell.fill = PatternFill("solid", fgColor=style[0])
        cell.font = Font(bold=True, color=style[1])
    cell.alignment = WRAP if wrap else TOP
    return cell


def table(ws, headers: list[str], rows: list[list[Any]], widths: list[int], *,
          styled: dict[int, dict[str, tuple[str, str]]] | None = None,
          wrapped: set[int] = frozenset()):
    for column, (header, width) in enumerate(zip(headers, widths), 1):
        cell = put(ws, 1, column, header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[get_column_letter(column)].width = width
    for row_number, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            styles = (styled or {}).get(column)
            put(ws, row_number, column, value, style=styles.get(value) if styles else None,
                wrap=column in wrapped)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"


def build_workbook(result: dict, design: Design | None = None) -> Workbook:
    if design is not None and result.get("input_sha256") is not None:
        current = hashlib.sha256(design.model_dump_json(exclude_none=False).encode()).hexdigest()
        if result["input_sha256"] != current:
            raise ValueError("check result does not match intermediate data")
    resources = {r.id: r for r in design.resources} if design else {}
    documents = {d.id: d.name for d in design.documents} if design else {}
    evidence = {e.id: e for e in design.evidence} if design else {}

    def evidence_text(ids: list[str]) -> str:
        lines = []
        for evidence_id in ids:
            e = evidence.get(evidence_id)
            if e is None:
                lines.append(evidence_id)
                continue
            where = (f"L{e.start_line}" if e.start_line == e.end_line
                     else f"L{e.start_line}-{e.end_line}")
            lines.append(f"{documents.get(e.document_id, e.document_id)} {where}: {e.excerpt}")
        return "\n".join(lines)

    def resource_cells(resource_id: str | None) -> list[str]:
        resource = resources.get(resource_id)
        return [resource.type if resource else "", resource.name if resource else "",
                resource_id or ""]

    wb = Workbook()
    _summary(wb.active, result, design)
    ws = wb.create_sheet("検査結果")
    findings = sorted(result.get("results", []),
                      key=lambda r: VERDICT_ORDER.index(r["verdict"])
                      if r["verdict"] in VERDICT_ORDER else len(VERDICT_ORDER))
    table(ws, ["No", "判定", "重大度", "ルール ID", "リソース種別", "リソース名", "リソース ID",
               "項目パス", "期待値", "実際の値", "理由", "根拠", "依存する未確定項目"],
          [[number, r["verdict"], r.get("severity"), r["rule_id"],
            *resource_cells(r.get("resource_id")), r.get("path") or "",
            text(r.get("expected")), text(r.get("actual")), r.get("reason", ""),
            evidence_text(r.get("evidence_ids", [])), "\n".join(map(text, r.get("dependencies", [])))]
           for number, r in enumerate(findings, 1)],
          [7, 18, 10, 40, 24, 16, 13, 34, 18, 18, 36, 56, 22],
          styled={2: VERDICT_STYLES}, wrapped={9, 10, 11, 12, 13})

    ws = wb.create_sheet("未検査範囲")
    table(ws, ["No", "種別", "リソース種別", "リソース名", "リソース ID", "項目パス", "理由"],
          [[number, item["kind"], *resource_cells(item.get("resource_id")),
            item.get("path") or "", item.get("reason", "")]
           for number, item in enumerate(result.get("coverage", []), 1)],
          [7, 24, 24, 16, 13, 34, 70], wrapped={7})

    if design is not None:
        ws = wb.create_sheet("設計値")
        rows = []
        for resource in design.resources:
            for field in resource.fields:
                selected = field.selected()
                if selected is not None:
                    shown, ids = text(selected.value), selected.evidence_ids
                else:
                    shown = "\n".join(text(c.value) for c in field.candidates)
                    ids = [i for c in field.candidates for i in c.evidence_ids]
                if field.default_intent:
                    shown = "(既定値に依存)" + (f"\n{shown}" if shown else "")
                    ids = ids + field.intent_evidence_ids
                rows.append([resource.type, resource.name, resource.id, field.path,
                             field.state.value, shown, evidence_text(ids)])
            for relation in design.relations:
                if relation.source_resource_id != resource.id:
                    continue
                target = resources.get(relation.target_resource_id)
                rows.append([resource.type, resource.name, resource.id, relation.source_path,
                             "REFERENCE" if target else "UNRESOLVED",
                             f"→ {target.type}/{target.name}" if target
                             else f"→ {relation.unresolved_name or ''}",
                             evidence_text(relation.evidence_ids)])
        table(ws, ["リソース種別", "リソース名", "リソース ID", "項目パス", "状態", "値", "根拠"],
              rows, [24, 16, 13, 34, 13, 34, 64],
              styled={5: STATE_STYLES}, wrapped={6, 7})
    return wb


def _summary(ws, result: dict, design: Design | None):
    ws.title = "概要"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 70
    put(ws, 1, 1, "AWS 設計チェック結果").font = TITLE_FONT
    row = 3
    status = result.get("status", "")
    items = [("処理状態", status), ("実行 ID", result.get("run_id", "")),
             ("生成日時", result.get("generated_at", "")),
             ("入力ハッシュ", result.get("input_sha256", ""))]
    if design is not None:
        items[1:1] = [("案件", design.project), ("環境", design.environment),
                      ("アカウント", design.account), ("リージョン", design.region)]
    if result.get("diagnostic"):
        items.append(("診断情報", result["diagnostic"]))
    status_style = VERDICT_STYLES["PASS" if status == "COMPLETE" else "FAIL"]
    for label, value in items:
        put(ws, row, 1, label).font = Font(bold=True)
        put(ws, row, 2, value, style=status_style if label == "処理状態" else None, wrap=True)
        row += 1

    row += 1
    put(ws, row, 1, "判定件数").font = SECTION_FONT
    row += 1
    summary = result.get("summary", {})
    for verdict in VERDICT_ORDER:
        put(ws, row, 1, verdict, style=VERDICT_STYLES[verdict])
        put(ws, row, 2, summary.get(verdict, 0)).alignment = Alignment(horizontal="left")
        row += 1
    put(ws, row, 1, "合計").font = Font(bold=True)
    put(ws, row, 2, sum(summary.values())).alignment = Alignment(horizontal="left")
    row += 1
    put(ws, row, 1, "未検査範囲").font = Font(bold=True)
    put(ws, row, 2, len(result.get("coverage", []))).alignment = Alignment(horizontal="left")

    row += 2
    put(ws, row, 1, "版").font = SECTION_FONT
    row += 1
    versions = dict(result.get("versions") or {})
    manifest = versions.pop("schema_manifest", None) or {}
    for key in ("region", "source", "retrieved_at", "zip_sha256"):
        if key in manifest:
            versions[f"schema_{key}"] = manifest[key]
    for key, value in versions.items():
        put(ws, row, 1, key)
        put(ws, row, 2, text(value), wrap=True)
        row += 1
    row += 1
    put(ws, row, 1, NOTE).font = Font(italic=True, color="595959")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export check results to an Excel workbook")
    parser.add_argument("result", type=Path, help="check result JSON")
    parser.add_argument("--design", type=Path,
                        help="intermediate design JSON; adds resource names, evidence and a design sheet")
    parser.add_argument("--output", required=True, type=Path, help="xlsx file to write")
    args = parser.parse_args(argv)
    try:
        result = json.loads(args.result.read_text(encoding="utf-8"))
        design = (Design.model_validate_json(args.design.read_text(encoding="utf-8"))
                  if args.design else None)
        workbook = build_workbook(result, design)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(args.output)
    except Exception as exc:
        print(f"FAILED: {exc}")
        return 2
    print(f"wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
