import copy
import json

import pytest
from openpyxl import load_workbook

from aws_design_sheet.document_cli import generate,main


def inputs(tmp_path):
    data={'schema_version':'1.1','system':{'id':'demo','name':'State reporting control'},
          'environments':[{'id':'prod','name':'Production','region':'ap-northeast-1'}],
          'resources':[
              {'id':'vpc','name':'VPC','type':'AWS::EC2::VPC','settings':{'CidrBlock':{'mode':'LITERAL','value':'10.0.0.0/16'}}},
              {'id':'subnet','name':'Subnet','type':'AWS::EC2::Subnet','settings':{'AvailabilityZoneId':{'status':'UNDECIDED'}}},
              {'id':'plan','name':'Plan','type':'AWS::Backup::BackupPlan','settings':{'BackupPlan':{'mode':'COMPOSITE','value':{
                  'BackupPlanName':'synthetic','BackupPlanRule':[{'RuleName':'daily','TargetBackupVault':{'$design':{
                      'mode':'EXPRESSION','expression':{'symbol':'vault.name','actual_value':None}}}}]}}}}]}
    trace=[]
    for sid,kind,status,item in [('SET-007','CFN','FICTIONAL_DECIDED','vpc.CidrBlock'),
                               ('SET-012','CFN','EXTERNAL_INPUT_REQUIRED','subnet.AvailabilityZoneId'),
                               ('SET-207','CFN','FICTIONAL_DECIDED','plan.BackupPlan'),
                               ('SET-005','APP','EXTERNAL_INPUT_REQUIRED',None)]:
        trace.append({'setting_id':sid,'kind':kind,'item_ids':[item] if item else [],'source_pointer':'/settings/'+sid,
                      'source_snapshot':{'setting_id':sid,'kind':kind,'status':status,'verification_status':'NOT_RUN'}})
    inp=tmp_path/'input.json';src=tmp_path/'trace.json'
    inp.write_text(json.dumps(data));src.write_text(json.dumps(trace))
    return inp,src,trace


def test_machine_report_keeps_original_states_and_supported_reference_reason(tmp_path):
    inp,src,trace=inputs(tmp_path);out=tmp_path/'generated'
    generate(inp,'prod',out,source_trace=src)
    report=json.loads((out/'review-report.json').read_text(encoding='utf-8'))
    assert report['source_status_counts']=={'EXTERNAL_INPUT_REQUIRED':2,'FICTIONAL_DECIDED':2}
    assert report['source_verification_counts']=={'NOT_RUN':4}
    assert report['normalized_status_counts']=={'DECIDED':2,'UNDECIDED':1}
    assert report['capabilities']['RESOURCE_REF']=={'supported':True,'physical_values_resolved':False}
    assert any(r['item_id']=='plan.BackupPlan' and r['reason_code']=='SYMBOL_TARGET_NOT_BOUND' for r in report['unresolved'])
    assert report['trace']['status']=='COMPLETE_FOR_GENERATED_INPUTS'
    assert report['counts']['source_kinds']=={'CFN':3,'APP':1}
    assert report['execution']['host_acceptance']=='NOT_VERIFIED_BY_CLI'
    assert report['execution']['aws_assessment']=='NOT_EVALUATED'
    assert json.loads((out/'design.json').read_text(encoding='utf-8'))['review_report']==report
    assert json.loads((out/'source-trace.json').read_text(encoding='utf-8'))==trace
    book=load_workbook(out/'design.xlsx',read_only=True)
    encoded=''.join(str(r[2] or '') for r in book['実行サマリ根拠'].iter_rows(min_row=2,values_only=True) if r[0]=='@review-report')
    assert json.loads(encoded)==report
    assert {r[0] for r in book['原資料状態'].iter_rows(min_row=2,values_only=True)}=={r['setting_id'] for r in trace}
    assert not any(c.data_type=='f' for sheet in book for row in sheet for c in row)
    book.close()


def test_without_trace_source_state_is_not_guessed_from_decided(tmp_path):
    inp,_,_=inputs(tmp_path);out=tmp_path/'generated';generate(inp,'prod',out)
    report=json.loads((out/'review-report.json').read_text(encoding='utf-8'))
    assert report['trace']['status']=='NOT_PROVIDED' and report['source_status_counts']=={}
    assert report['normalized_status_counts']['DECIDED']==2
    assert report['execution']['whole_design_approval']=='NOT_ESTABLISHED'


@pytest.mark.parametrize('mutation',['duplicate','missing_state','bad_binding','app_binding','mismatched_id'])
def test_malformed_trace_fails_before_outputs(tmp_path,mutation):
    inp,src,trace=inputs(tmp_path)
    if mutation=='duplicate':trace.append(copy.deepcopy(trace[0]))
    elif mutation=='missing_state':trace[0]['source_snapshot'].pop('status')
    elif mutation=='bad_binding':trace[0]['item_ids']=['does-not-exist']
    elif mutation=='app_binding':trace[-1]['item_ids']=['vpc.CidrBlock']
    elif mutation=='mismatched_id':trace[0]['source_snapshot']['setting_id']='OTHER'
    src.write_text(json.dumps(trace));out=tmp_path/'failed'
    assert main([str(inp),'--environment','prod','--output',str(out),'--source-trace',str(src)])==2
    assert not out.exists()


def test_failed_evaluator_is_reported_as_failure_not_host_success(tmp_path,monkeypatch):
    from aws_design_sheet import document_evaluation
    def fail(_):raise RuntimeError('synthetic evaluator failure')
    monkeypatch.setattr(document_evaluation,'run_checker',fail)
    inp,src,_=inputs(tmp_path);out=tmp_path/'failed-evaluator'
    assert main([str(inp),'--environment','prod','--output',str(out),'--evaluate','--source-trace',str(src)])==3
    report=json.loads((out/'review-report.json').read_text(encoding='utf-8'))
    assert report['execution']['evaluation_status']=='FAILED' and report['execution']['aws_assessment']=='ERROR'
    assert report['execution']['host_acceptance']=='NOT_VERIFIED_BY_CLI'
    assert report['evaluation_stages']['checker']['status']=='FAILED'
