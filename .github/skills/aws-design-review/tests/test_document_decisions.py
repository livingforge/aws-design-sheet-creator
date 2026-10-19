"""Basic-design decisions: outline, explicit traceability and generated documents."""
import copy
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from aws_design_sheet.document_cli import generate
from aws_design_sheet.document_normalize import normalize
from aws_design_sheet.document_v11 import SchemaInput

EXAMPLE=Path(__file__).resolve().parents[1]/'examples/basic-design/input.json'
def sample():return json.loads(EXAMPLE.read_text(encoding='utf-8'))
def trace(data,env='prod'):return normalize(json.dumps(data).encode(),env).decision_trace
def codes(report):return {f['finding_id'] for f in report['findings']}
def decision(report,ident):return next(d for d in report['decisions'] if d['decision_id']==ident)
def resource(data,ident):return next(r for r in data['resources'] if r['id']==ident)
def rows(book,sheet):return [r for r in book[sheet].iter_rows(min_row=2,values_only=True) if str(r[-1]).startswith('1/')]


def test_sample_generates_trace_workbook_and_basic_design(tmp_path):
    out=tmp_path/'prod';summary=generate(EXAMPLE,'prod',out)
    report=json.loads((out/'decision-trace.json').read_text(encoding='utf-8'))
    assert summary['design_decision_status']=='COMPLETE' and summary['design_decision_assessment']=='GAPS_FOUND'
    assert summary['design_decisions']==report['summary'] and summary['basic_design_file']=='basic-design.md'
    assert {d['decision_id']:d['realization'] for d in report['decisions']}=={
        'dec-overview':'DOCUMENT_ONLY','dec-vpc-cidr':'PARTIAL','dec-encryption':'REALIZED','dec-versioning':'REALIZED',
        'dec-log-retention':'REALIZED','dec-edge':'NOT_APPLICABLE','dec-backup':'UNDECIDED','dec-contact':'DOCUMENT_ONLY'}
    assert decision(report,'dec-versioning')['mechanization']=='CHECKED'
    assert decision(report,'dec-versioning')['policy_results']=={'NEEDS_REVIEW':1}
    assert codes(report)=={'REALIZING_SETTING_UNDECIDED:dec-vpc-cidr:subnet-c',
        'DECISION_UNDECIDED:dec-backup','SECTION_EMPTY:non-functional','SECTION_EMPTY:account','SECTION_EMPTY:compute',
        'SETTING_WITHOUT_BASIS:vpc.EnableDnsSupport','SETTING_WITHOUT_BASIS:app-logs.LogGroupName'}
    # Composite leaves keep their own link and pointer.
    assert {l['path'] for l in report['links'] if l['decision_id']=='dec-encryption'}=={
        '/properties/BucketEncryption','/properties/BucketEncryption/ServerSideEncryptionConfiguration/0/ServerSideEncryptionByDefault/KMSMasterKeyID',
        '/properties/KmsKeyId'}
    design=json.loads((out/'design.json').read_text(encoding='utf-8'))
    assert design['decision_trace']==report
    assert design['review_report']['execution']['design_decision_assessment']=='GAPS_FOUND'
    book=load_workbook(out/'design.xlsx',read_only=True)
    assert {r[0] for r in rows(book,'設計方針')}=={d['decision_id'] for d in report['decisions']}
    assert {r[0] for r in rows(book,'方針指摘')}==codes(report)
    assert len({(r[0],r[1]) for r in rows(book,'方針トレース')})==len(report['links'])
    assert not any(c.data_type=='f' for sheet in book for row in sheet for c in row)
    book.close()
    part=load_workbook(out/'services-01.xlsx',read_only=True)
    retention=''.join(r[6] or '' for r in part['Logs設定'].iter_rows(min_row=2,values_only=True) if r[0]=='app-logs.RetentionInDays')
    assert '方針: dec-log-retention' in retention and '設計方針' not in part.sheetnames
    part.close()
    text=(out/'basic-design.md').read_text(encoding='utf-8')
    for expected in ('# 架空ポータル 基本設計書','## 4. ネットワーク','### 4.1 VPCとサブネットのアドレス','## 2. 非機能要件\n\n**未記載**',
                     '| app-logs.RetentionInDays | /properties/RetentionInDays | 決定済み / INTERNAL / 明示値<br>400 | RESOLVED | 本番は監査対応で13か月保持 |',
                     '```mermaid','## 付録B 未決事項・方針トレースの指摘','## 付録C 出典'):
        assert expected in text


def test_environment_override_and_environment_scope():
    data=sample()
    dev=trace(data,'dev')
    assert decision(dev,'dec-log-retention')['realization']=='REALIZED'
    assert next(l for l in dev['links'] if l['item_id']=='app-logs.RetentionInDays')['rationale']==''
    data['decisions'][4]['applies_to']['environments']=['prod']
    assert decision(trace(data,'dev'),'dec-log-retention')['realization']=='NOT_IN_ENVIRONMENT'
    assert decision(trace(data,'prod'),'dec-log-retention')['realization']=='REALIZED'


def test_state_and_scope_conflicts_are_reported():
    data=sample()
    resource(data,'app-logs')['settings']['LogGroupName']['decision_ids']=['dec-backup']
    resource(data,'vpc')['settings']['EnableDnsSupport']['decision_ids']=['dec-edge']
    resource(data,'app-logs')['settings']['RetentionInDays']['decision_ids'].append('dec-versioning')
    report=trace(data,'dev')
    assert report['assessment']=='CONFLICTS_FOUND'
    assert {'DECISION_UNDECIDED_BUT_SETTING_DECIDED:dec-backup:app-logs.LogGroupName:/properties/LogGroupName',
            'DECISION_NOT_APPLICABLE_BUT_LINKED:dec-edge:vpc.EnableDnsSupport:/properties/EnableDnsSupport',
            'LINK_OUTSIDE_DECISION_SCOPE:dec-versioning:app-logs.RetentionInDays:/properties/RetentionInDays'}<=codes(report)


def test_fully_traced_input_has_no_gaps():
    data=sample()
    data['decisions']=[d for d in data['decisions'] if d['id'] not in ('dec-backup',)]
    data['requirements']=[r for r in data['requirements'] if r['id']!='req-avail-01']
    data['design_outline']=[{'id':s,'title':s} for s in ('overview','network','data','edge','security','monitoring','operations')]
    resource(data,'subnet-c')['settings']['CidrBlock']={'mode':'LITERAL','value':'10.0.3.0/24','decision_ids':['dec-vpc-cidr']}
    resource(data,'vpc')['settings']['EnableDnsSupport']['rationale']='名前解決のため'
    resource(data,'app-logs')['settings']['LogGroupName']['sources']=['src-req']
    report=trace(data)
    assert report['assessment']=='NO_GAPS_WITHIN_DECLARED_SCOPE' and report['findings']==[]
    assert report['outline_source']=='DECLARED'


def test_optional_outline_section_is_not_a_gap():
    data=sample();data['design_outline']=[{'id':'overview','title':'概要'},{'id':'network','title':'NW'},{'id':'data','title':'データ'},
        {'id':'edge','title':'公開'},{'id':'security','title':'セキュリティ'},{'id':'monitoring','title':'監視'},
        {'id':'backup-dr','title':'DR'},{'id':'operations','title':'運用'},{'id':'compute','title':'計算','required':False}]
    report=trace(data)
    assert not any(f['code']=='SECTION_EMPTY' for f in report['findings'])
    assert [s['section_id'] for s in report['sections']][-1]=='compute'


@pytest.mark.parametrize('mutation',['unknown_link','nested_unknown_link','duplicate_link','unknown_section','unknown_requirement',
    'unknown_policy','unknown_source','scope_resource','scope_type','scope_environment','empty_scope','duplicate_id','duplicate_section'])
def test_invalid_declarations_are_rejected(mutation):
    data=sample();d=data['decisions'][1]
    if mutation=='unknown_link':resource(data,'vpc')['settings']['CidrBlock']['decision_ids']=['dec-missing']
    elif mutation=='nested_unknown_link':
        leaf=resource(data,'assets')['settings']['BucketEncryption']['value']['ServerSideEncryptionConfiguration'][0]
        leaf['ServerSideEncryptionByDefault']['KMSMasterKeyID']['$design']['decision_ids']=['dec-missing']
    elif mutation=='duplicate_link':resource(data,'vpc')['settings']['CidrBlock']['decision_ids']=['dec-vpc-cidr']*2
    elif mutation=='unknown_section':d['section']='not-in-outline'
    elif mutation=='unknown_requirement':d['requirement_ids']=['req-missing']
    elif mutation=='unknown_policy':d['policy_ids']=['pol-missing']
    elif mutation=='unknown_source':d['sources']=['src-missing']
    elif mutation=='scope_resource':d['applies_to']['resources']=['missing']
    elif mutation=='scope_type':d['applies_to']={'types':['AWS::Fictional::Thing']}
    elif mutation=='scope_environment':d['applies_to']['environments']=['stg']
    elif mutation=='empty_scope':d['applies_to']={'environments':['prod']}
    elif mutation=='duplicate_id':d['id']='vpc'
    elif mutation=='duplicate_section':data['design_outline']=[{'id':'network','title':'a'},{'id':'network','title':'b'}]
    with pytest.raises(ValidationError):SchemaInput.model_validate(data)


def test_input_without_decisions_is_not_checked_and_still_renders(tmp_path):
    data=sample()
    for key in ('decisions','project_policies'):data.pop(key)
    for r in data['resources']:
        for s in r['settings'].values():s.pop('decision_ids',None)
    leaf=resource(data,'assets')['settings']['BucketEncryption']
    leaf.pop('decision_ids',None)
    leaf['value']['ServerSideEncryptionConfiguration'][0]['ServerSideEncryptionByDefault']['KMSMasterKeyID']['$design'].pop('decision_ids')
    data['environments'][1]['overrides'][0]['settings']['RetentionInDays'].pop('decision_ids')
    path=tmp_path/'input.json';path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
    summary=generate(path,'prod',tmp_path/'out')
    assert summary['design_decision_status']=='NOT_CHECKED' and 'decision_trace_file' not in summary
    assert not (tmp_path/'out/decision-trace.json').exists()
    book=load_workbook(tmp_path/'out/design.xlsx',read_only=True)
    assert '設計方針' not in book.sheetnames;book.close()
    assert '## 4. ネットワーク\n\n**未記載**' in (tmp_path/'out/basic-design.md').read_text(encoding='utf-8')


def test_markdown_cells_escape_table_syntax(tmp_path):
    data=sample();data['decisions'][1]['alternatives']=['a|b']
    resource(data,'vpc')['settings']['CidrBlock']['rationale']='x | y\nz'
    path=tmp_path/'input.json';path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
    generate(path,'prod',tmp_path/'out')
    text=(tmp_path/'out/basic-design.md').read_text(encoding='utf-8')
    assert '| x \\| y<br>z |' in text and '  - a|b' in text


def decision_input(data,ident):return next(d for d in data['decisions'] if d['id']==ident)


def test_expansion_fills_absent_setting_and_is_traced(tmp_path):
    out=tmp_path/'prod';generate(EXAMPLE,'prod',out)
    design=json.loads((out/'design.json').read_text(encoding='utf-8'))
    item=next(i for i in design['items'] if i['item_id']=='app-logs.KmsKeyId')
    assert item['origin']=='decision' and item['input_pointer']=='/decisions/2/expansions/0'
    assert item['expected']['decision_ids']==['dec-encryption'] and item['expected']['mode']=='EXTERNAL_REF'
    report=design['decision_trace']
    assert {(r['item_id'],r['result']) for r in report['expansions']}=={('app-logs.KmsKeyId','APPLIED'),('assets.VersioningConfiguration','MATCHED_EXPLICIT')}
    assert decision(report,'dec-encryption')['realization']=='REALIZED'
    assert report['summary']['expansions']=={'APPLIED':1,'MATCHED_EXPLICIT':1}
    book=load_workbook(out/'design.xlsx',read_only=True)
    assert {r[0] for r in rows(book,'方針展開')}=={'dec-encryption/0','dec-versioning/0'};book.close()
    part=load_workbook(out/'services-01.xlsx',read_only=True)
    text=''.join(r[6] or '' for r in part['Logs設定'].iter_rows(min_row=2,values_only=True) if r[0]=='app-logs.KmsKeyId')
    assert '出所: 方針から展開' in text;part.close()
    assert '方針 dec-encryption から展開：ログ用のカスタマー管理キー' in (out/'basic-design.md').read_text(encoding='utf-8')


def test_explicit_setting_wins_and_difference_is_a_conflict():
    data=sample()
    decision_input(data,'dec-log-retention')['expansions']=[{'property':'RetentionInDays','mode':'LITERAL','value':30}]
    dev,prod=trace(data,'dev'),trace(data,'prod')
    assert [r['result'] for r in dev['expansions'] if r['decision_id']=='dec-log-retention']==['MATCHED_EXPLICIT']
    assert [r['result'] for r in prod['expansions'] if r['decision_id']=='dec-log-retention']==['OVERRIDDEN']
    assert 'EXPANSION_OVERRIDDEN:dec-log-retention:app-logs.RetentionInDays' in codes(prod)
    document=normalize(json.dumps(data).encode(),'prod')
    assert next(i for i in document.items if i.item_id=='app-logs.RetentionInDays').expected.value==400
    decision_input(data,'dec-log-retention')['expansions'][0]['environments']=['dev']
    assert not [r for r in trace(data,'prod')['expansions'] if r['decision_id']=='dec-log-retention']


def test_expansion_feeds_required_input_checklist_and_skips_undecided():
    data=sample();resource(data,'subnet-c')['settings'].pop('VpcId')
    data['decisions'].append({'id':'dec-subnet-vpc','section':'network','title':'所属VPC','statement':'サブネットはポータルVPCに置く。',
        'applies_to':{'types':['AWS::EC2::Subnet']},'expansions':[{'property':'VpcId','mode':'RESOURCE_REF','ref':'vpc'}]})
    document=normalize(json.dumps(data).encode(),'prod')
    row=next(r for r in document.checklist.rows if r['resource_id']=='subnet-c' and r['path']=='/properties/VpcId')
    assert row['result']=='SATISFIED_REFERENCE' and row['input_pointer']=='/decisions/8/expansions/0'
    assert {r['result'] for r in document.decision_trace['expansions'] if r['decision_id']=='dec-subnet-vpc'}=={'APPLIED','MATCHED_EXPLICIT'}
    decision_input(data,'dec-subnet-vpc')['status']='UNDECIDED'
    document=normalize(json.dumps(data).encode(),'prod')
    assert not any(i.item_id=='subnet-c.VpcId' for i in document.items if i.origin=='decision')
    assert next(r for r in document.checklist.rows if r['resource_id']=='subnet-c' and r['path']=='/properties/VpcId')['result']=='MISSING'


@pytest.mark.parametrize('mutation',['no_scope','unwritable','wrong_type','secret_property','secret_class','ambiguous',
    'bad_environment','bad_type','missing_value','bad_reference'])
def test_invalid_expansions_are_rejected(mutation):
    data=sample();d=decision_input(data,'dec-encryption');e=d['expansions'][0]
    if mutation=='no_scope':
        data['decisions'].append({'id':'dec-x','section':'operations','title':'x','statement':'x',
                                  'expansions':[{'property':'LogGroupName','mode':'LITERAL','value':'x'}]})
    elif mutation=='unwritable':e['property']='Arn'
    elif mutation=='wrong_type':e['types']=[]  # KmsKeyId is not a writable S3 bucket property
    elif mutation=='secret_property':
        data['resources'].append({'id':'param','name':'p','type':'AWS::SSM::Parameter'})
        data['decisions'].append({'id':'dec-x','section':'operations','title':'x','statement':'x','applies_to':{'resources':['param']},
                                  'expansions':[{'property':'Value','mode':'LITERAL','value':'plain'}]})
    elif mutation=='secret_class':e['classification']='SECRET'
    elif mutation=='ambiguous':
        data['decisions'].append({'id':'dec-x','section':'operations','title':'x','statement':'x','applies_to':{'resources':['app-logs']},
                                  'expansions':[{'property':'KmsKeyId','mode':'EXTERNAL_REF','ref':'other'}]})
    elif mutation=='bad_environment':e['environments']=['stg']
    elif mutation=='bad_type':e['types']=['AWS::Fictional::Thing']
    elif mutation=='missing_value':e.update(mode='LITERAL',ref=None)
    elif mutation=='bad_reference':e.update(mode='RESOURCE_REF',ref='missing')
    with pytest.raises(ValidationError):SchemaInput.model_validate(data)
