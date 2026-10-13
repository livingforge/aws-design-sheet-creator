"""Explicit project constraints stay independent from AWS validity and equality."""
import copy
import json
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from aws_design_sheet.document_cli import generate, main
from aws_design_sheet.document_normalize import normalize
from aws_design_sheet.document_policy_model import PolicyConfigurationError


def lit(value): return {'mode': 'LITERAL', 'value': value, 'sources': ['source']}


def spec(kind='AWS::S3::Bucket', settings=None, path='/properties/PublicAccessBlockConfiguration/BlockPublicAcls', values=None):
    return {'schema_version': '1.1', 'system': {'id': 'demo', 'name': 'Project policy'},
        'environments': [{'id': 'dev', 'name': 'Dev', 'region': 'ap-northeast-1'},
                         {'id': 'prod', 'name': 'Prod', 'region': 'ap-northeast-1'}],
        'resources': [{'id': 'resource', 'name': 'Resource', 'cfn_logical_id': 'Resource', 'type': kind,
                      'settings': settings if settings is not None else {'PublicAccessBlockConfiguration': lit({'BlockPublicAcls': True})}}],
        'requirements': [{'id': 'req', 'title': 'Project requirement', 'text': 'Human declared scope'}],
        'sources': [{'id': 'source', 'title': 'Signed design decision', 'locator': 'spec:section-1'}],
        'project_policies': [{'id': 'rule', 'requirement_id': 'req', 'resource_id': 'resource', 'path': path,
            'values': [True] if values is None else values, 'sources': ['source'], 'environments': ['prod'],
            'confirmation': 'CONFIRMED', 'rationale': 'Explicit project choice, not an AWS-wide prohibition'}]}


def norm(data, env='prod', cfn=None):
    return normalize(json.dumps(data).encode(), env, json.dumps(cfn).encode() if cfn else None)


def policy(data, **kwargs): return norm(data, **kwargs).project_policy


@pytest.mark.parametrize('actual,expected,result', [(True, True, 'PASS'), (False, True, 'FAIL'),
    (False, 0, 'FAIL'), (True, 1, 'FAIL'), (False, False, 'PASS')])
def test_boolean_is_not_number(actual, expected, result):
    data = spec(settings={'PublicAccessBlockConfiguration': lit({'BlockPublicAcls': actual})}, values=[expected])
    assert policy(data).rows[0]['result'] == result


@pytest.mark.parametrize('value,wanted,result', [(0, 0, 'PASS'), (0, False, 'FAIL'), (1, 0, 'FAIL')])
def test_zero_is_a_concrete_number(value, wanted, result):
    data = spec('AWS::SQS::Queue', {'DelaySeconds': lit(value)}, '/properties/DelaySeconds', [wanted])
    assert policy(data).rows[0]['result'] == result


@pytest.mark.parametrize('value,wanted,result', [('', '', 'PASS'), ('', None, 'FAIL'), ('0', 0, 'FAIL'), ('false', False, 'FAIL')])
def test_empty_string_is_distinct_and_requires_explicit_expectation(value, wanted, result):
    data = spec('AWS::SQS::Queue', {'QueueName': lit(value)}, '/properties/QueueName', [wanted])
    assert policy(data).rows[0]['result'] == result


def test_explicit_null_differs_from_missing():
    data=spec('AWS::StepFunctions::StateMachine',{'Definition':lit({'Comment':None})},'/properties/Definition/Comment',[None])
    row=policy(data).rows[0]
    assert row['result']=='PASS' and row['trace']['actual_available'] and row['trace']['actual'] is None
    data['resources'][0]['settings']['Definition']=lit({})
    assert policy(data).rows[0]['result']=='NEEDS_REVIEW'


def test_aws_allowed_runtime_can_violate_project_policy_and_equal_cfn(tmp_path):
    data = spec('AWS::Lambda::Function', {'Runtime': lit('nodejs22.x'), 'Handler': lit('index.handler'),
        'Role': lit('arn:aws:iam::123456789012:role/Role'), 'Code': lit({'ZipFile': 'exports.handler=async()=>{};'})},
        '/properties/Runtime', ['python3.13'])
    data['project_policies'][0]['operator'] = 'IN'
    cfn = {'Resources': {'Resource': {'Type': 'AWS::Lambda::Function', 'Properties': {k:v['value'] for k,v in data['resources'][0]['settings'].items()}}}}
    doc = norm(data, cfn=cfn)
    assert next(i for i in doc.items if i.name == 'Runtime').comparison == 'MATCH'
    assert doc.project_policy.rows[0]['result'] == 'FAIL'
    path = tmp_path/'input.json'; path.write_text(json.dumps(data))
    summary = generate(path, 'prod', tmp_path/'out', evaluate_design=True)
    assert summary['generation_status'] == 'GENERATED' and summary['evaluation_status'] == 'COMPLETE'
    assert summary['project_policy_assessment'] == 'NONCOMPLIANT'
    aws = json.loads((tmp_path/'out/evaluation.json').read_text(encoding='utf-8'))
    assert not any(r['path'] == '/properties/Runtime' and r['verdict'] == 'FAIL' for r in aws['results'])
    data['resources'][0]['settings']['Runtime'] = lit('python3.13')
    assert policy(data).rows[0]['result'] == 'PASS'


@pytest.mark.parametrize('setting', [None, {'status':'UNDECIDED'}, {'status':'NOT_APPLICABLE'}, {'mode':'OMITTED'},
    {'mode':'DEFAULT'}, {'mode':'EXTERNAL_REF','ref':'external://value'},
    {'mode':'EXPRESSION','expression':{'Ref':'External'}},
    {'mode':'EXPRESSION','expression':{'Fn::GetAtt':['External','Arn']}},
    {'mode':'SECRET_REF','classification':'SECRET','ref':'secret://opaque'}])
def test_missing_and_unknown_cannot_pass_or_borrow_cfn(setting):
    data = spec('AWS::SQS::Queue', {'FifoQueue': setting} if setting else {}, '/properties/FifoQueue')
    cfn = {'Resources':{'Resource':{'Type':'AWS::SQS::Queue','Properties':{'FifoQueue':True}}}}
    result = policy(data, cfn=cfn)
    assert result.rows[0]['result'] == 'NEEDS_REVIEW'
    assert not result.rows[0]['trace']['actual_available']


@pytest.mark.parametrize('condition,expected,verdict', [(True,'TRUE','PASS'),(False,'FALSE','NOT_APPLICABLE'),(None,'UNKNOWN','NEEDS_REVIEW')])
def test_condition_true_false_unknown(condition, expected, verdict):
    data = spec('AWS::SQS::Queue', {'FifoQueue':lit(True)}, '/properties/FifoQueue')
    if condition is not None: data['resources'][0]['settings']['ContentBasedDeduplication'] = lit(condition)
    data['project_policies'][0]['when'] = [{'path':'/properties/ContentBasedDeduplication','values':[True]}]
    result = policy(data).rows[0]
    assert (result['condition_result'],result['result']) == (expected,verdict)


def test_conditions_are_bounded_conjunction_and_membership():
    data=spec('AWS::SQS::Queue',{'FifoQueue':lit(True),'QueueName':lit('project.fifo'),'DelaySeconds':lit(0)},'/properties/FifoQueue')
    data['project_policies'][0]['when']=[{'path':'/properties/QueueName','operator':'IN','values':['project.fifo','other.fifo']},
                                      {'path':'/properties/DelaySeconds','values':[0]}]
    assert policy(data).rows[0]['result']=='PASS'
    data['resources'][0]['settings']['DelaySeconds']=lit(1)
    assert policy(data).rows[0]['result']=='NOT_APPLICABLE'


def array_spec(values):
    return spec('AWS::EC2::SecurityGroup', {'GroupDescription':lit('Project'), 'SecurityGroupIngress':
        {'mode':'COMPOSITE','value':[{'IpProtocol':v} for v in values]}}, '/properties/SecurityGroupIngress/*/IpProtocol', ['tcp'])


def test_array_all_members_nested_decisions_and_empty_selection():
    data=array_spec(['tcp','udp',{'$design':{'status':'UNDECIDED','sources':['source']}}])
    result=policy(data)
    assert [r['result'] for r in result.rows]==['PASS','FAIL','NEEDS_REVIEW']
    assert len({r['policy_check_id'] for r in result.rows})==3
    assert result.rows[2]['trace']['input_pointer'].endswith('/value/2/IpProtocol/$design')
    assert policy(array_spec([])).rows[0]['result']=='NEEDS_REVIEW'
    explicit=copy.deepcopy(data);explicit['project_policies'][0]['path']='/properties/SecurityGroupIngress/1/IpProtocol'
    assert policy(explicit).rows[0]['result']=='FAIL'


def test_environment_override_and_scope_keep_original_identity():
    data=spec();data['environments'][1]['overrides']=[{'resource_id':'resource','settings':{
        'PublicAccessBlockConfiguration':lit({'BlockPublicAcls':False})}}]
    row=policy(data).rows[0]
    assert row['result']=='FAIL' and row['trace']['input_pointer'].startswith('/environments/1/overrides/0/')
    assert row['trace']['row_id']=='prod/resource.PublicAccessBlockConfiguration'
    assert policy(data,env='dev').rows[0]['result']=='NOT_APPLICABLE'


def exception(data, confirmation='CONFIRMED'):
    data['policy_exceptions']=[{'id':'waiver','policy_id':'rule','paths':[data['project_policies'][0]['path']],
        'environments':['prod'],'sources':['source'],'reason':'Explicit temporary project exception', 'confirmation':confirmation}]


def test_exception_has_scope_and_is_never_pass_or_unknown_waiver():
    data=spec(settings={'PublicAccessBlockConfiguration':lit({'BlockPublicAcls':False})})
    exception(data)
    result=policy(data)
    assert result.rows[0]['result']=='EXEMPTED' and result.assessment=='REVIEW'
    data['policy_exceptions'][0]['confirmation']='UNCONFIRMED'
    assert policy(data).rows[0]['result']=='FAIL'
    data['resources'][0]['settings']={};data['policy_exceptions'][0]['confirmation']='CONFIRMED'
    assert policy(data).rows[0]['result']=='NEEDS_REVIEW'


def test_unregistered_requirements_and_unconfirmed_rule_do_not_imply_coverage():
    data=spec();data['requirements'].append({'id':'unregistered','title':'Other requirement','text':'Needs manual mapping'})
    result=policy(data)
    assert result.rows[0]['result']=='PASS' and result.assessment=='REVIEW'
    assert result.summary['unregistered_requirements']==['unregistered']
    assert result.provenance['all_requirements_evaluated'] is False
    data['project_policies'][0]['confirmation']='UNCONFIRMED'
    assert policy(data).rows[0]['result']=='NEEDS_REVIEW'
    data['project_policies']=[]
    assert policy(data).status=='NOT_REGISTERED'


@pytest.mark.parametrize('field,value', [('resource_id','unknown'),('requirement_id','unknown'),('sources',['unknown']),
    ('environments',['unknown']),('path','/properties/NotAProperty')])
def test_bad_references_are_configuration_errors(field,value):
    data=spec();data['project_policies'][0][field]=value
    with pytest.raises(PolicyConfigurationError):norm(data)


@pytest.mark.parametrize('patch', [{'operator':'EXEC','values':['SECRET_CANARY_CONFIG']},
    {'path':'/properties/~bad'},{'values':[]},{'values':[True,False]},{'environments':[]},{'rationale':'  '},
    {'execute':'SECRET_CANARY_CONFIG'}])
def test_bad_declaration_syntax_cli_is_safe(tmp_path,capsys,patch):
    data=spec();data['project_policies'][0].update(patch)
    path=tmp_path/'input.json';path.write_text(json.dumps(data))
    assert main([str(path),'--environment','prod','--output',str(tmp_path/'out')])==2
    output=capsys.readouterr().out
    assert json.loads(output)['generation_status']=='POLICY_CONFIGURATION_ERROR'
    assert 'SECRET_CANARY_CONFIG' not in output and not (tmp_path/'out').exists()


def test_wildcards_and_exceptions_cannot_broaden_scope():
    data=spec('AWS::Lambda::Function',{'Environment':lit({'Variables':{'PUBLIC':'public'}})},'/properties/Environment/Variables/*',['public'])
    with pytest.raises(PolicyConfigurationError):norm(data)
    data=spec();exception(data);data['policy_exceptions'][0]['paths']=['/properties/PublicAccessBlockConfiguration/IgnorePublicAcls']
    with pytest.raises(PolicyConfigurationError):norm(data)
    data=spec();exception(data);data['policy_exceptions'].append({**data['policy_exceptions'][0],'id':'waiver-two'})
    with pytest.raises(PolicyConfigurationError):norm(data)


def test_secret_comparand_rejected_without_logging_value(tmp_path,capsys):
    data=spec('AWS::Lambda::Function',{'Environment':{'mode':'COMPOSITE','value':{'Variables':{'OPAQUE':{
        '$design':{'mode':'SECRET_REF','classification':'SECRET','ref':'secret://opaque'}}}}}},
        '/properties/Environment/Variables/OPAQUE',['SECRET_CANARY_EXPECTATION'])
    path=tmp_path/'input.json';path.write_text(json.dumps(data))
    assert main([str(path),'--environment','prod','--output',str(tmp_path/'out')])==2
    output=capsys.readouterr().out
    assert 'SECRET_CANARY_EXPECTATION' not in output
    assert json.loads(output)['error_code']=='SECRET_POLICY_EXPECTATION'


def read_policy_book(path):
    book=load_workbook(path,data_only=False);parts={}
    for row in book['案件方針根拠'].iter_rows(min_row=2,values_only=True):parts[row[0]]=parts.get(row[0],'')+(row[2] or '')
    decoded={k:json.loads(v) for k,v in parts.items()};result=decoded.pop('@project-policy');result['rows']=list(decoded.values())
    for row in book['案件方針'].iter_rows(min_row=2):
        assert row[0].hyperlink
        target=int(row[0].hyperlink.target.rsplit('A',1)[1])
        assert book['案件方針根拠'].cell(target,1).value==row[0].value
    assert all(c.data_type!='f' for sheet in book for row in sheet for c in row)
    book.close();return result


def test_json_excel_roundtrip_and_cfn_secret_canary(tmp_path):
    data=array_spec(['tcp','udp']);data['project_policies'][0]['rationale']='=FORMULA()\n'+('Long evidence / '*40)
    source=tmp_path/'input.json';source.write_text(json.dumps(data))
    cfn=tmp_path/'cfn.json';cfn.write_text(json.dumps({'Resources':{'Other':{'Type':'AWS::IAM::User','Properties':{'LoginProfile':{'Password':'SECRET_CANARY_CFN'}}}}}))
    folder=tmp_path/'out';generate(source,'prod',folder,cfn)
    expected=json.loads((folder/'project-policy.json').read_text(encoding='utf-8'))
    assert read_policy_book(folder/'design.xlsx')==expected
    for p in folder.iterdir():
        if p.suffix=='.json':assert 'SECRET_CANARY_CFN' not in p.read_text(encoding='utf-8')
        if p.suffix=='.xlsx':
            with ZipFile(p) as z:assert all(b'SECRET_CANARY_CFN' not in z.read(n) for n in z.namelist())


def test_evaluator_failure_and_selection_limit_never_pass(monkeypatch,tmp_path,capsys):
    from aws_design_sheet.document_policy import evaluate
    document=norm(array_spec(['tcp','tcp']))
    monkeypatch.setattr('aws_design_sheet.document_policy.MAX_SELECTIONS',1)
    assert evaluate(document).assessment=='ERROR'
    def fail(*args):raise RuntimeError('SECRET_CANARY_EXCEPTION')
    monkeypatch.setattr('aws_design_sheet.document_policy.select',fail)
    data=spec();path=tmp_path/'input.json';path.write_text(json.dumps(data))
    assert main([str(path),'--environment','prod','--output',str(tmp_path/'out')])==3
    assert 'SECRET_CANARY_EXCEPTION' not in capsys.readouterr().out
    assert json.loads((tmp_path/'out/project-policy.json').read_text(encoding='utf-8'))['assessment']=='ERROR'


def test_version_one_input_policy_is_supported():
    data=spec('AWS::EC2::VPC',{'EnableDnsSupport':lit(True)},'/properties/EnableDnsSupport')
    data['schema_version']='1.0'
    assert policy(data).rows[0]['result']=='PASS'
