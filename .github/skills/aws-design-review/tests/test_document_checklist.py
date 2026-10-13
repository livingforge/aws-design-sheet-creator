"""Completeness is independent of CFN equality, input presence and generation."""
import copy
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from aws_design_sheet.document_catalog import catalog
from aws_design_sheet.document_checklist import Inspector, Unknown, predicate
from aws_design_sheet.document_cli import generate
from aws_design_sheet.document_normalize import normalize
from aws_design_sheet.document_requirements import InputRequirement
from aws_design_sheet.document_v11 import SchemaSetting


ROOT = Path(__file__).resolve().parents[1]


def spec(kind='AWS::Lambda::Function', settings=None):
    return {'schema_version': '1.1', 'system': {'id': 'demo', 'name': 'Checklist'},
            'environments': [{'id': 'dev', 'name': 'Development', 'region': 'ap-northeast-1'},
                             {'id': 'prod', 'name': 'Production', 'region': 'ap-northeast-1'}],
            'resources': [{'id': 'resource', 'name': 'Resource', 'type': kind,
                           'cfn_logical_id': 'Resource', 'settings': settings or {}}]}


def norm(data, cfn=None, env='dev'):
    return normalize(json.dumps(data).encode(), env, json.dumps(cfn).encode() if cfn else None)


def rows(data, path):
    return [r for r in data.checklist.rows if r['path'] == path]


def synthetic(schema, settings, requirements=()):
    cat = copy.copy(catalog())
    kind = 'AWS::SQS::Queue'
    cat.schemas = {kind: schema}
    resource = SimpleNamespace(id='r', type=kind, settings={k: SchemaSetting.model_validate(v) for k, v in settings.items()})
    env = SimpleNamespace(id='dev', overrides=[])
    doc = SimpleNamespace(environment_id='dev', input_sha256='test', design=SimpleNamespace(
        resources=[resource], environments=[env], sources=[], required_inputs=list(requirements)))
    return Inspector(doc, cat, include_extensions=False).run()


def lit(value): return {'mode': 'LITERAL', 'value': value}


def test_missing_on_both_sides_and_generation_still_succeeds(tmp_path):
    data = spec()
    cfn = {'Resources': {'Resource': {'Type': 'AWS::Lambda::Function', 'Properties': {}}}}
    doc = norm(data, cfn)
    assert not doc.items
    assert rows(doc, '/properties/Role')[0]['result'] == 'MISSING'
    assert rows(doc, '/properties/Code')[0]['result'] == 'MISSING'
    path = tmp_path / 'input.json'; path.write_text(json.dumps(data))
    summary = generate(path, 'dev', tmp_path / 'output')
    assert summary['generation_status'] == 'GENERATED'
    assert summary['input_completeness'] == 'INCOMPLETE'
    assert summary['evaluation_status'] == 'NOT_EVALUATED'


@pytest.mark.parametrize('setting,expected,state', [
    (None, 'MISSING', 'MISSING'), (lit(''), 'NEEDS_REVIEW', 'EMPTY'), (lit(None), 'NEEDS_REVIEW', 'NULL'),
    (lit([]), 'NEEDS_REVIEW', 'EMPTY'), (lit({}), 'NEEDS_REVIEW', 'EMPTY'),
    (lit(False), 'SATISFIED', 'VALUE'), (lit(0), 'SATISFIED', 'VALUE'),
    ({'status': 'UNDECIDED'}, 'NEEDS_REVIEW', 'UNDECIDED'),
    ({'status': 'NOT_APPLICABLE'}, 'MISSING', 'NOT_APPLICABLE'),
    ({'mode': 'OMITTED'}, 'MISSING', 'OMITTED'),
    ({'mode': 'DEFAULT'}, 'NEEDS_REVIEW', 'DEFAULT_REQUESTED'),
    ({'mode': 'EXTERNAL_REF', 'ref': 'external://role'}, 'NEEDS_REVIEW', 'EXTERNAL_REFERENCE'),
    ({'mode': 'EXPRESSION', 'expression': {'Ref': 'External'}}, 'NEEDS_REVIEW', 'EXPRESSION_UNRESOLVED'),
    ({'mode': 'SECRET_REF', 'classification': 'SECRET', 'ref': 'secret://test'}, 'NEEDS_REVIEW', 'SECRET_REFERENCE'),
    ({'mode': 'RESOURCE_REF', 'ref': 'other'}, 'SATISFIED_REFERENCE', 'INTERNAL_REFERENCE')])
def test_presence_is_not_sufficiency(setting, expected, state):
    result = synthetic({'properties': {'Needed': {}}, 'required': ['Needed']}, {'Needed': setting} if setting else {})
    row = result.rows[0]
    assert (row['result'], row['input_state'], row['necessity']) == (expected, state, 'REQUIRED')


@pytest.mark.parametrize('value,condition,result', [(True, 'TRUE', 'MISSING'), (False, 'FALSE', 'NOT_APPLICABLE'),
                                                 (Unknown('EXTERNAL_REFERENCE'), 'UNKNOWN', 'NEEDS_REVIEW')])
def test_if_and_dependencies_three_values(value, condition, result):
    schema = {'properties': {'Enabled': {}, 'Needed': {}},
              'if': {'properties': {'Enabled': {'const': True}}, 'required': ['Enabled']},
              'then': {'required': ['Needed']}}
    setting = {'mode': 'EXTERNAL_REF', 'ref': 'external://flag'} if isinstance(value, Unknown) else lit(value)
    row = synthetic(schema, {'Enabled': setting}).rows[0]
    assert (row['condition_result'], row['result']) == (condition, result)
    schema = {'properties': {'Enabled': {}, 'Needed': {}}, 'dependencies': {'Enabled': ['Needed']}}
    result = synthetic(schema, {'Enabled': {'status': 'UNDECIDED'}})
    assert result.rows[0]['result'] == 'NEEDS_REVIEW'
    assert result.rows[0]['schema_pointer'] == '#/dependencies/Enabled/0'


def test_nested_required_arrays_and_readonly():
    schema = {'properties': {'Objects': {'type': 'array', 'items': {'type': 'object',
        'properties': {'Name': {'type': 'string'}, 'Computed': {}}, 'required': ['Name', 'Computed']}}},
        'readOnlyProperties': ['/properties/Objects/*/Computed']}
    result = synthetic(schema, {'Objects': lit([{'Name': 'ok'}, {}])})
    assert [(r['path'], r['result']) for r in result.rows] == [
        ('/properties/Objects/0/Name', 'SATISFIED'), ('/properties/Objects/1/Name', 'MISSING')]
    assert result.summary['read_only_excluded'] == 2
    unknown = synthetic(schema, {'Objects': {'mode': 'EXTERNAL_REF', 'ref': 'external://list'}})
    assert unknown.rows[0]['path'] == '/properties/Objects/*/Name'
    assert unknown.rows[0]['result'] == 'NEEDS_REVIEW'


def test_oneof_alternatives_are_never_flattened_into_all_required():
    schema = {'properties': {'A': {}, 'B': {}}, 'oneOf': [{'required': ['A']}, {'required': ['B']}]}
    good = synthetic(schema, {'A': lit('chosen')})
    assert good.rows[0]['result'] == 'SATISFIED'
    assert next(r for r in good.rows if r['path'].endswith('/B'))['result'] == 'NOT_APPLICABLE'
    missing = synthetic(schema, {})
    assert missing.rows[0]['result'] == 'BRANCH_CONFLICT'
    assert all(r['necessity'] == 'UNDETERMINED' for r in missing.rows[1:])
    conflict = synthetic(schema, {'A': lit('a'), 'B': lit('b')})
    assert conflict.rows[0]['result'] == 'BRANCH_CONFLICT'
    unresolved = synthetic(schema, {'A': {'status': 'UNDECIDED'}})
    assert unresolved.rows[0]['result'] == 'NEEDS_REVIEW'


def test_predicates_do_not_invent_values_or_ignore_unknown_keywords():
    assert predicate({'type': 'object'}, {'Opaque': Unknown('EXTERNAL_REFERENCE')}) is True
    assert predicate({'type': 'array'}, [Unknown('EXTERNAL_REFERENCE')]) is True
    assert predicate({'properties': {'Flag': {'const': False}}, 'required': ['Flag']}, {'Flag': False}) is True
    assert predicate({'properties': {'Flag': {'const': False}}, 'required': ['Flag']}, {}) is False
    assert predicate({'properties': {'Flag': {'const': False}}}, {'Flag': Unknown('DEFAULT_REQUESTED')}) is None
    assert predicate({'futureConstraint': 'opaque'}, {}) is None
    assert predicate({'$ref': '#/definitions/X', 'definitions': {'X': {'required': ['N']}}}, {}) is False


def project_spec():
    data = spec('AWS::Logs::LogGroup')
    data['sources'] = [{'id': 'source', 'title': 'Retention decision', 'locator': 'minutes#decision-1'}]
    data['requirements'] = [{'id': 'retention', 'title': 'Retention', 'text': 'Set retention explicitly.'}]
    data['required_inputs'] = [{'id': 'retention-input', 'requirement_id': 'retention', 'resource_id': 'resource',
        'path': '/properties/RetentionInDays', 'sources': ['source'], 'confirmation': 'CONFIRMED', 'environments': ['prod']}]
    return data


def test_project_requirement_environment_missing_source_and_repair():
    data = project_spec()
    assert norm(data).checklist.rows[-1]['result'] == 'NOT_APPLICABLE'
    assert norm(data, env='prod').checklist.rows[-1]['result'] == 'MISSING'
    data['resources'][0]['settings']['RetentionInDays'] = lit(30)
    assert norm(data, env='prod').checklist.rows[-1]['result'] == 'SATISFIED'
    data['required_inputs'][0]['sources'] = []
    row = norm(data, env='prod').checklist.rows[-1]
    assert row['result'] == 'NEEDS_REVIEW' and row['evidence'] == [{'source': 'UNKNOWN'}]


def test_explicit_default_requires_schema_evidence_and_project_permission():
    data = project_spec()
    data['resources'][0]['type'] = 'AWS::SQS::Queue'
    data['required_inputs'][0]['path'] = '/properties/VisibilityTimeout'
    data['resources'][0]['settings'] = {'VisibilityTimeout': {'mode': 'DEFAULT'}}
    row = norm(data, env='prod').checklist.rows[-1]
    # The pinned schema is the authority: absence of default metadata is not a default.
    has_default = 'default' in catalog().schemas['AWS::SQS::Queue']['properties']['VisibilityTimeout']
    assert row['result'] == 'NEEDS_REVIEW'
    data['required_inputs'][0]['allow_documented_default'] = True
    assert norm(data, env='prod').checklist.rows[-1]['result'] == ('SATISFIED_DEFAULT' if has_default else 'NEEDS_REVIEW')


def test_cfn_lint_conditional_requirement_reused_and_cfn_cannot_fill_input():
    data = spec(settings={'Code': lit({'ZipFile': 'def handler(event, context): return None'})})
    doc = norm(data)
    runtime = rows(doc, '/properties/Runtime')
    assert any(r['result'] == 'MISSING' and r['evidence'][0]['provider'] == 'cfn-lint' for r in runtime)
    cfn = {'Resources': {'Resource': {'Type': 'AWS::Lambda::Function', 'Properties': {'Runtime': 'python3.12'}}}}
    assert any(r['result'] == 'MISSING' for r in rows(norm(data, cfn), '/properties/Runtime'))
    data['resources'][0]['settings']['Runtime'] = lit('python3.12')
    assert all(r['result'] == 'SATISFIED' for r in rows(norm(data), '/properties/Runtime'))


def test_checklist_json_and_xlsx_roundtrip_long_evidence(tmp_path):
    data = project_spec()
    data['required_inputs'][0]['rationale'] = '=not-a-formula\n日本語🙂' * 5000
    inp = tmp_path / 'input.json'; inp.write_text(json.dumps(data, ensure_ascii=False), encoding='utf8')
    out = tmp_path / 'generated'; generate(inp, 'prod', out)
    expected = json.loads((out / 'checklist.json').read_text(encoding='utf8'))
    model = json.loads((out / 'design.json').read_text(encoding='utf8'))
    assert model['checklist'] == expected
    book = load_workbook(out / 'design.xlsx', read_only=True, data_only=False)
    collected = {}
    for row in book['必要入力'].iter_rows(min_row=2):
        values = [c.value or '' for c in row]
        assert all(c.data_type != 'f' for c in row)
        collected.setdefault(values[0], []).append(values[3])
    assert set(collected) == {r['checklist_id'] for r in expected['rows']}
    parts = [str(row[2] or '') for row in book['必要入力根拠'].iter_rows(min_row=2, values_only=True)
             if row[0] == expected['rows'][-1]['checklist_id']]
    details = json.loads(''.join(parts))
    assert details['detail'] == data['required_inputs'][0]['rationale']
    book.close()


def test_v1_input_remains_supported():
    data = spec('AWS::EC2::Subnet'); data['schema_version'] = '1.0'
    doc = norm(data)
    assert doc.model_version == '1.0'
    assert rows(doc, '/properties/VpcId')[0]['result'] == 'MISSING'


def test_readonly_is_never_required_and_cannot_decide_a_condition():
    schema = {'properties': {'Output': {}, 'Input': {}}, 'readOnlyProperties': ['/properties/Output'],
              'oneOf': [{'required': ['Output']}],
              'if': {'required': ['Output']}, 'then': {'required': ['Input']}}
    result = synthetic(schema, {})
    assert len(result.rows) == 1 and result.rows[0]['path'] == '/properties/Input'
    assert result.rows[0]['condition_result'] == 'UNKNOWN'
    assert result.rows[0]['result'] == 'NEEDS_REVIEW'


def test_anyof_refs_dependencies_else_and_unsupported_condition():
    schema = {'properties': {'A': {}, 'B': {}, 'Mode': {}},
              'definitions': {'A': {'required': ['A']}, 'B': {'required': ['B']}},
              'anyOf': [{'$ref': '#/definitions/A'}, {'$ref': '#/definitions/B'}]}
    result = synthetic(schema, {'B': lit(0)})
    assert result.status == 'SATISFIED_WITHIN_SCOPE'
    assert any(r['path'].endswith('/A') and r['result'] == 'NOT_APPLICABLE' for r in result.rows)
    conditional = {'properties': schema['properties'], 'if': {'properties': {'Mode': {'const': 'first'}}, 'required': ['Mode']},
                   'then': {'required': ['A']}, 'else': {'required': ['B']}}
    result = synthetic(conditional, {'Mode': lit('second'), 'B': lit(False)})
    assert [r['result'] for r in result.rows] == ['NOT_APPLICABLE', 'SATISFIED']
    conditional['if'] = {'futureUnknownKeyword': {}}
    assert all(r['result'] == 'NEEDS_REVIEW' for r in synthetic(conditional, {}).rows)


def test_explicit_project_default_does_not_waive_aws_required():
    from aws_design_sheet.document_model import Source
    schema = {'properties': {'Needed': {'type': 'integer', 'default': 0}}, 'required': ['Needed']}
    cat = copy.copy(catalog()); cat.schemas = {'AWS::SQS::Queue': schema}
    resource = SimpleNamespace(id='r', type='AWS::SQS::Queue', settings={'Needed': SchemaSetting(mode='DEFAULT')})
    req = InputRequirement(id='project', requirement_id='q', resource_id='r', path='/properties/Needed',
                           sources=['source'], confirmation='CONFIRMED', allow_documented_default=True)
    doc = SimpleNamespace(environment_id='dev', input_sha256='test', design=SimpleNamespace(resources=[resource],
        environments=[SimpleNamespace(id='dev', overrides=[])], required_inputs=[req],
        sources=[Source(id='source', title='Approved default policy', locator='example#policy')]))
    result = Inspector(doc, cat, include_extensions=False).run()
    assert [r['result'] for r in result.rows] == ['NEEDS_REVIEW', 'SATISFIED_DEFAULT']
    assert result.status == 'NEEDS_REVIEW'
    assert result.rows[-1]['default_evidence'][0]['schema_pointer'] == '#/properties/Needed/default'
    assert result.rows[-1]['value_resolution'] == 'UNRESOLVED'


def test_environment_override_and_nested_unknown_sources():
    data = spec('AWS::S3::Bucket')
    data['sources'] = [{'id': 'leaf', 'title': 'Leaf decision', 'locator': 'source#line'}]
    data['resources'][0]['settings']['Tags'] = {'mode': 'COMPOSITE', 'value': [
        {'Key': 'Owner', 'Value': {'$design': {'status': 'UNDECIDED', 'sources': ['leaf']}}}]}
    doc = norm(data)
    value = rows(doc, '/properties/Tags/0/Value')[0]
    assert value['result'] == 'NEEDS_REVIEW' and value['input_sources'][0]['id'] == 'leaf'
    data['environments'][1]['overrides'] = [{'resource_id': 'resource', 'settings': {'Tags': lit([{'Key': 'Owner', 'Value': 'team'}])}}]
    fixed = rows(norm(data, env='prod'), '/properties/Tags/0/Value')[0]
    assert fixed['result'] == 'SATISFIED'
    assert fixed['input_pointer'].startswith('/environments/1/overrides/0/')


def test_declared_omission_can_match_cfn_but_not_waive_required_input():
    data = spec(settings={'Role': {'mode': 'OMITTED'}})
    doc = norm(data, {'Resources': {'Resource': {'Type': 'AWS::Lambda::Function', 'Properties': {}}}})
    assert next(i for i in doc.items if i.name == 'Role').comparison == 'MATCH'
    assert 'OMISSION' in next(i for i in doc.items if i.name == 'Role').comparison_basis
    assert rows(doc, '/properties/Role')[0]['result'] == 'MISSING'


def test_project_readonly_unknown_paths_and_secret_predicates_rejected():
    for path in ['/properties/Arn', '/properties/Nonexistent']:
        data = project_spec(); data['required_inputs'][0]['path'] = path
        with pytest.raises(ValueError): norm(data)
    data = spec('AWS::RDS::DBInstance')
    data['requirements'] = [{'id': 'r', 'title': 'Secret policy', 'text': 'No values in predicates.'}]
    data['required_inputs'] = [{'id': 'p', 'requirement_id': 'r', 'resource_id': 'resource', 'path': '/properties/Engine',
                               'when': {'properties': {'MasterUserPassword': {'const': 'fictional-value'}}}}]
    with pytest.raises(ValueError): norm(data)
    data = project_spec(); data['required_inputs'][0]['when'] = {'properties': 'invalid'}
    with pytest.raises(ValueError): norm(data)


def test_unknown_tuple_and_recursive_schema_are_bounded():
    schema = {'properties': {'Items': {'type': 'array', 'items': [{'properties': {'A': {}}, 'required': ['A']}],
                                      'additionalItems': False}}}
    result = synthetic(schema, {'Items': {'mode': 'EXTERNAL_REF', 'ref': 'external://array'}})
    assert result.rows[0]['result'] == 'NEEDS_REVIEW'
    recursive = {'properties': {'A': {}}, '$ref': '#/definitions/Loop', 'definitions': {'Loop': {'$ref': '#/definitions/Loop'}}}
    result = synthetic(recursive, {})
    assert result.status == 'NEEDS_REVIEW' and len(result.rows) == 1


def test_custom_cfn_lint_preprocessing_is_not_applied_at_wrong_path():
    doc = norm(spec('AWS::StepFunctions::StateMachine'))
    assert not rows(doc, '/properties/StartAt')
    assert not rows(doc, '/properties/States')
    custom = [r for r in doc.checklist.rows if r['evidence'][0].get('rule_id') == 'E3601']
    assert len(custom) == 1 and custom[0]['result'] == 'NEEDS_REVIEW'
    assert custom[0]['input_state'] == 'NOT_INSPECTED'


def test_secret_omission_has_no_payload_and_does_not_fetch_values():
    data = spec('AWS::RDS::DBInstance', {'MasterUserPassword': {'mode': 'OMITTED', 'classification': 'SECRET'}})
    doc = norm(data)
    assert doc.items[0].expected.mode == 'OMITTED'
    data['resources'][0]['settings']['MasterUserPassword']['value'] = 'fictional-marker'
    with pytest.raises(ValueError): norm(data)


def test_v1_different_region_does_not_claim_pinned_aws_requirements_apply():
    data = spec('AWS::EC2::Subnet'); data['schema_version'] = '1.0'
    data['environments'][0]['region'] = 'eu-west-1'
    doc = norm(data)
    assert doc.generation_status == 'GENERATED'
    assert rows(doc, '/properties/VpcId')[0]['result'] == 'NEEDS_REVIEW'
    assert doc.checklist.summary['missing_or_conflicting'] == 0
    assert doc.checklist.status == 'NEEDS_REVIEW'


def test_shared_required_definition_keeps_different_conditions_and_stable_ids():
    schema = {'properties': {'First': {}, 'Second': {}, 'Needed': {}},
        'definitions': {'Required': {'required': ['Needed']}},
        'allOf': [{'if': {'properties': {flag: {'const': True}}, 'required': [flag]},
                   'then': {'$ref': '#/definitions/Required'}} for flag in ('First', 'Second')]}
    settings = {'First': lit(False), 'Second': lit(True)}
    incomplete = synthetic(schema, settings)
    assert [r['result'] for r in incomplete.rows] == ['NOT_APPLICABLE', 'MISSING']
    repaired = synthetic(schema, {**settings, 'Needed': lit('supplied')})
    assert [r['result'] for r in repaired.rows] == ['NOT_APPLICABLE', 'SATISFIED']
    assert [r['checklist_id'] for r in incomplete.rows] == [r['checklist_id'] for r in repaired.rows]
